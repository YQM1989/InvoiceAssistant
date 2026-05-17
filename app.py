from __future__ import annotations

import csv
import os
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import font as tkfont
from tkinter import filedialog, messagebox, simpledialog, ttk

from openpyxl import Workbook, load_workbook

from invoice_parser import parse_invoice, sanitize_filename
from storage import (
    InvoiceRecord,
    InvoiceStore,
    WhitelistEntry,
    clean_company_name,
    is_valid_company_name,
    normalize_company,
    normalize_invoice_number,
)


APP_TITLE = "发票助手"
DEFAULT_PAGE_SIZE = 50
PAGE_SIZE_OPTIONS = (50, 100)
AMOUNT_REVIEW_THRESHOLD = 100000
APPROVAL_STATUS_LABELS = {
    "": "全部",
    "normal": "正常",
    "pending": "待审批",
    "review": "需复核",
}


def configure_tk_runtime() -> None:
    candidates = [
        Path(r"C:\Users\Public\codex_tcl"),
        Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent)),
        Path(getattr(sys, "_MEIPASS", Path(sys.executable).resolve().parent)) / "_internal",
        Path(sys.executable).resolve().parent,
        Path(sys.executable).resolve().parent / "_internal",
        Path(sys.base_prefix) / "tcl",
    ]
    for candidate in candidates:
        if (candidate / "_tcl_data").exists() or (candidate / "_tk_data").exists():
            tcl_dir = candidate / "_tcl_data"
            tk_dir = candidate / "_tk_data"
        else:
            tcl_dir = candidate / "tcl8.6"
            tk_dir = candidate / "tk8.6"
        if tcl_dir.exists():
            os.environ["TCL_LIBRARY"] = str(tcl_dir)
        if tk_dir.exists():
            os.environ["TK_LIBRARY"] = str(tk_dir)
        if tcl_dir.exists() or tk_dir.exists():
            break

class InvoiceAssistantApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title(APP_TITLE)
        self.geometry("1520x920")
        self.minsize(1320, 760)

        self.font_family = self._choose_ui_font()
        self.option_add("*Font", (self.font_family, 10))
        self.configure(bg="#F3F5F8")

        self.base_dir = Path(__file__).resolve().parent
        self.runtime_root = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else self.base_dir
        self.data_root = self._resolve_data_root()
        self._migrate_legacy_runtime_data()
        self.invoice_root_dir = self.data_root / "invoice_files"
        self.approved_invoice_dir = self.invoice_root_dir / "approved"
        self.pending_invoice_dir = self.invoice_root_dir / "pending"
        self.export_dir = self.data_root / "exports"
        self.approved_invoice_dir.mkdir(parents=True, exist_ok=True)
        self.pending_invoice_dir.mkdir(parents=True, exist_ok=True)
        self.export_dir.mkdir(parents=True, exist_ok=True)

        self.store = InvoiceStore(self.data_root / "invoice_assistant.db")

        self.current_page = 1
        self.total_items = 0
        self.sort_column = "date"
        self.sort_ascending = False
        self.var_page_size = tk.IntVar(value=DEFAULT_PAGE_SIZE)
        self.month_picker_popup: tk.Toplevel | None = None
        self.month_picker_year = datetime.now().year
        self.invoice_heading_labels = {
            "id": "ID",
            "seller": "销售方",
            "buyer": "购买方",
            "invoice_no": "发票号码",
            "amount": "金额",
            "date": "开票日期",
            "file": "文件名",
            "imported_at": "导入时间",
        }

        self.invoice_seller_options: list[str] = []
        self.invoice_buyer_options: list[str] = []
        self.month_options: list[str] = []
        self.pending_seller_options: list[str] = []
        self.pending_buyer_options: list[str] = []
        self.reconcile_seller_options: list[str] = []
        self.reconcile_buyer_options: list[str] = []

        self.var_filter_seller = tk.StringVar()
        self.var_filter_buyer = tk.StringVar()
        self.var_start_month = tk.StringVar()
        self.var_end_month = tk.StringVar()
        self.var_filter_keyword = tk.StringVar()
        self.var_status = tk.StringVar(value="就绪")
        self.var_total_amount = tk.StringVar(value="¥ 0.00")

        self.var_pending_seller = tk.StringVar()
        self.var_pending_buyer = tk.StringVar()
        self.var_pending_invoice_no = tk.StringVar()
        self.var_pending_amount = tk.StringVar()
        self.var_pending_date = tk.StringVar()
        self.var_pending_reason = tk.StringVar()
        self.var_pending_review_reason = tk.StringVar()
        self.var_pending_file = tk.StringVar()
        self.var_pending_seller_sync = tk.StringVar(value="none")
        self.var_pending_buyer_sync = tk.StringVar(value="none")

        self.var_rec_seller = tk.StringVar()
        self.var_rec_buyer = tk.StringVar()
        self.var_rec_no = tk.StringVar()
        self.var_rec_amount = tk.StringVar()
        self.reconcile_rows: list[dict[str, object]] = []
        self.reconcile_sort_column = "date"
        self.reconcile_sort_ascending = False
        self.reconcile_heading_labels = {
            "seller": "销售方",
            "buyer": "购买方",
            "invoice_no": "发票号",
            "amount": "金额",
            "date": "开票日期",
            "status": "结果",
            "file": "文件名",
        }
        self.var_batch_reconcile_only_issues = tk.BooleanVar(value=False)
        self.batch_reconcile_results: list[dict[str, str]] = []
        self.batch_reconcile_sort_column = "matched_date"
        self.batch_reconcile_sort_ascending = False
        self.batch_reconcile_heading_labels = {
            "input_seller": "输入销售方",
            "input_buyer": "输入购买方",
            "input_invoice_no": "输入发票号",
            "input_amount": "输入金额",
            "result": "核对结果",
            "matched_seller": "系统销售方",
            "matched_buyer": "系统购买方",
            "matched_amount": "系统金额",
            "matched_date": "开票日期",
            "file": "文件名",
        }

        self.whitelist_widgets: dict[str, dict[str, object]] = {}
        self.whitelist_search_vars = {
            "seller": tk.StringVar(),
            "buyer": tk.StringVar(),
        }
        self.whitelist_alias_search_vars = {
            "seller": tk.StringVar(),
            "buyer": tk.StringVar(),
        }
        self.month_picker_target = "start"

        self._build_style()
        self._build_ui()
        self.refresh_all()

    def _choose_ui_font(self) -> str:
        available = set(tkfont.families(self))
        for family in ("Microsoft YaHei UI", "Microsoft YaHei", "Segoe UI", "Arial Unicode MS", "Arial"):
            if family in available:
                return family
        return "TkDefaultFont"

    def _resolve_data_root(self) -> Path:
        if getattr(sys, "frozen", False):
            data_root = self.runtime_root / "_data"
        else:
            data_root = self.base_dir
        data_root.mkdir(parents=True, exist_ok=True)
        return data_root

    def _migrate_legacy_runtime_data(self) -> None:
        return

    def _build_style(self) -> None:
        style = ttk.Style(self)
        style.theme_use("clam")
        base_bg = "#F5F5F7"
        panel_bg = "#FFFFFF"
        panel_soft = "#FAFAFC"
        graphite = "#24364B"
        graphite_hover = "#1C2A3A"
        accent = "#24364B"
        accent_hover = "#1C2A3A"
        selected = "#375C7C"
        border = "#D8DEE8"
        soft_border = "#E6EAF0"
        heading_bg = "#EEF3F8"
        button_bg = "#F2F3F5"
        button_active = "#E8EAED"
        text = "#1D1D1F"
        muted = "#6E6E73"
        font = self.font_family
        self.configure(bg=base_bg)

        style.configure("TFrame", background=base_bg)
        style.configure("White.TFrame", background=panel_bg)
        style.configure("Card.TFrame", background=panel_bg)
        style.configure("Panel.TFrame", background=panel_bg, bordercolor=soft_border, relief="solid", borderwidth=1)
        style.configure("Toolbar.TFrame", background=panel_bg)
        style.configure("TLabel", background=base_bg, foreground=text)
        style.configure("White.TLabel", background=panel_bg, foreground=text)
        style.configure("Muted.TLabel", background=panel_bg, foreground=muted)
        style.configure("SectionTitle.TLabel", background=panel_bg, foreground="#4B5563", font=(font, 10, "bold"))
        style.configure("FormLabel.TLabel", background=panel_bg, foreground="#334155", font=(font, 10))
        style.configure("Footer.TLabel", background=base_bg, foreground="#9CA3AF", font=(font, 9))

        style.configure("TNotebook", background=base_bg, borderwidth=0, tabmargins=(0, 0, 0, 0))
        style.layout(
            "TNotebook.Tab",
            [
                (
                    "Notebook.tab",
                    {
                        "sticky": "nswe",
                        "children": [
                            (
                                "Notebook.padding",
                                {
                                    "side": "top",
                                    "sticky": "nswe",
                                    "children": [("Notebook.label", {"side": "top", "sticky": ""})],
                                },
                            )
                        ],
                    },
                )
            ],
        )
        style.configure(
            "TNotebook.Tab",
            padding=(18, 9),
            font=(font, 10),
            background="#F2F4F7",
            foreground="#3A3A3C",
            borderwidth=0,
            focuscolor=base_bg,
        )
        style.map(
            "TNotebook.Tab",
            background=[("selected", graphite), ("active", "#E7EBF0")],
            foreground=[("selected", "white"), ("active", text)],
            padding=[("selected", (22, 10))],
        )

        style.configure(
            "Treeview",
            rowheight=30,
            font=(font, 10),
            background=panel_bg,
            fieldbackground=panel_bg,
            foreground=text,
            bordercolor=soft_border,
            lightcolor=soft_border,
            darkcolor=soft_border,
        )
        style.map("Treeview", background=[("selected", selected)], foreground=[("selected", "white")])
        style.configure(
            "Treeview.Heading",
            font=(font, 10, "bold"),
            background=heading_bg,
            foreground=text,
            relief="flat",
            padding=(8, 8),
        )

        button_layout = [
            (
                "Button.border",
                {
                    "sticky": "nswe",
                    "border": "1",
                    "children": [
                        (
                            "Button.padding",
                            {
                                "sticky": "nswe",
                                "children": [("Button.label", {"sticky": "nswe"})],
                            },
                        )
                    ],
                },
            )
        ]
        style.layout("TButton", button_layout)
        style.layout("Primary.TButton", button_layout)
        style.configure("TButton", padding=(12, 7), font=(font, 10), background=button_bg, foreground=text, bordercolor=border, focusthickness=0, focuscolor=button_bg)
        style.map("TButton", background=[("active", button_active), ("pressed", "#DDE1E7")])
        style.configure("Primary.TButton", padding=(15, 8), font=(font, 10, "bold"), background=accent, foreground="white", bordercolor=accent, focusthickness=0, focuscolor=accent)
        style.map("Primary.TButton", background=[("active", accent_hover), ("pressed", graphite_hover)], foreground=[("active", "white"), ("pressed", "white")])

        style.configure("TEntry", padding=(6, 4), fieldbackground=panel_bg, foreground=text, bordercolor=border, lightcolor=border, darkcolor=border, insertcolor=text)
        style.configure("TCombobox", padding=(6, 4), fieldbackground=panel_bg, background=panel_bg, foreground=text, bordercolor=border, arrowcolor=muted)
        style.map("TCombobox", fieldbackground=[("readonly", panel_bg)], selectbackground=[("readonly", panel_bg)], selectforeground=[("readonly", text)])
        style.configure("TCheckbutton", background=base_bg, foreground=text, font=(font, 10))
        style.configure("TRadiobutton", background=base_bg, foreground=text, font=(font, 10))
        style.configure("TScrollbar", background="#E5E7EB", troughcolor=panel_soft, bordercolor=panel_soft, arrowcolor=muted)
        style.configure("TPanedwindow", background=base_bg)

        style.configure("Card.TLabelframe", background=panel_bg, bordercolor=soft_border, relief="solid")
        style.configure("Card.TLabelframe.Label", background=base_bg, foreground="#3A3A3C", font=(font, 10, "bold"))
        style.configure("SummaryLabel.TLabel", background=panel_bg, foreground=muted)
        style.configure("SummaryValue.TLabel", background=panel_bg, foreground=graphite, font=(font, 18, "bold"))

    def _build_ui(self) -> None:
        container = ttk.Frame(self, padding=12)
        container.pack(fill=tk.BOTH, expand=True)

        notebook = ttk.Notebook(container)
        notebook.pack(fill=tk.BOTH, expand=True)

        self.tab_invoice = ttk.Frame(notebook)
        self.tab_approval = ttk.Frame(notebook)
        self.tab_whitelist = ttk.Frame(notebook)
        self.tab_reconcile = ttk.Frame(notebook)

        notebook.add(self.tab_invoice, text="发票管理")
        notebook.add(self.tab_approval, text="审批池")
        notebook.add(self.tab_whitelist, text="白名单")
        notebook.add(self.tab_reconcile, text="对账")

        self._build_invoice_tab()
        self._build_approval_tab()
        self._build_whitelist_tab()
        self._build_reconcile_tab()

        footer = ttk.Frame(container)
        footer.pack(fill=tk.X, pady=(8, 0))
        footer.columnconfigure(0, weight=1)
        footer.columnconfigure(1, weight=0, minsize=86)
        ttk.Label(footer, textvariable=self.var_status, anchor="w", style="Footer.TLabel").grid(row=0, column=0, sticky="ew")
        ttk.Label(footer, text="by YQM", anchor="e", style="Footer.TLabel").grid(row=0, column=1, sticky="e", padx=(12, 4))

    def _build_invoice_tab(self) -> None:
        top = ttk.Frame(self.tab_invoice, padding=(12, 10), style="Panel.TFrame")
        top.pack(fill=tk.X, padx=6, pady=(6, 8))

        import_group = ttk.Frame(top, style="Toolbar.TFrame")
        import_group.pack(side=tk.LEFT, padx=(0, 24))
        ttk.Label(import_group, text="导入", style="SectionTitle.TLabel").pack(anchor="w", pady=(0, 6))
        import_buttons = ttk.Frame(import_group, style="Toolbar.TFrame")
        import_buttons.pack(anchor="w")
        ttk.Button(import_buttons, text="导入PDF文件", command=self.import_files, style="Primary.TButton", takefocus=0).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(import_buttons, text="导入PDF文件夹", command=self.import_folder, takefocus=0).pack(side=tk.LEFT)

        export_group = ttk.Frame(top, style="Toolbar.TFrame")
        export_group.pack(side=tk.LEFT, padx=(0, 24))
        ttk.Label(export_group, text="导出", style="SectionTitle.TLabel").pack(anchor="w", pady=(0, 6))
        export_buttons = ttk.Frame(export_group, style="Toolbar.TFrame")
        export_buttons.pack(anchor="w")
        ttk.Button(export_buttons, text="导出汇总表", command=self.export_summary_excel, takefocus=0).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(export_buttons, text="导出PDF文件", command=self.export_filtered_pdfs, takefocus=0).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Button(export_buttons, text="分目录导出PDF(月份_销售方)", command=self.export_filtered_grouped, takefocus=0).pack(side=tk.LEFT)

        misc_group = ttk.Frame(top, style="Toolbar.TFrame")
        misc_group.pack(side=tk.LEFT)
        ttk.Label(misc_group, text="操作", style="SectionTitle.TLabel").pack(anchor="w", pady=(0, 6))
        misc_buttons = ttk.Frame(misc_group, style="Toolbar.TFrame")
        misc_buttons.pack(anchor="w")
        ttk.Button(misc_buttons, text="刷新", command=self.refresh_all, takefocus=0).pack(side=tk.LEFT)

        filter_frame = ttk.Frame(self.tab_invoice, padding=(14, 12), style="Panel.TFrame")
        filter_frame.pack(fill=tk.X, padx=6, pady=(0, 8))
        for idx in range(7):
            filter_frame.columnconfigure(idx, weight=0)
        filter_frame.columnconfigure(1, weight=1)
        filter_frame.columnconfigure(3, weight=1)
        filter_frame.columnconfigure(5, weight=1)

        ttk.Label(filter_frame, text="筛选", style="SectionTitle.TLabel").grid(row=0, column=0, columnspan=7, sticky="w", pady=(0, 10))

        ttk.Label(filter_frame, text="销售方", style="FormLabel.TLabel").grid(row=1, column=0, sticky="w", padx=(0, 8))
        self.cbo_filter_seller = ttk.Combobox(filter_frame, textvariable=self.var_filter_seller, width=22)
        self.cbo_filter_seller.grid(row=1, column=1, sticky="ew", padx=(0, 18))
        self.cbo_filter_seller.bind("<KeyRelease>", lambda _e: self.filter_combobox_values(self.cbo_filter_seller, self.invoice_seller_options, self.var_filter_seller.get()))
        self.cbo_filter_seller.bind("<<ComboboxSelected>>", lambda _e: self.apply_filter(show_no_result_feedback=True))

        ttk.Label(filter_frame, text="购买方", style="FormLabel.TLabel").grid(row=1, column=2, sticky="w", padx=(0, 8))
        self.cbo_filter_buyer = ttk.Combobox(filter_frame, textvariable=self.var_filter_buyer, width=22)
        self.cbo_filter_buyer.grid(row=1, column=3, sticky="ew", padx=(0, 18))
        self.cbo_filter_buyer.bind("<KeyRelease>", lambda _e: self.filter_combobox_values(self.cbo_filter_buyer, self.invoice_buyer_options, self.var_filter_buyer.get()))
        self.cbo_filter_buyer.bind("<<ComboboxSelected>>", lambda _e: self.apply_filter(show_no_result_feedback=True))

        ttk.Label(filter_frame, text="开票月份", style="FormLabel.TLabel").grid(row=2, column=0, sticky="w", padx=(0, 8), pady=(10, 0))
        start_month_frame = ttk.Frame(filter_frame)
        start_month_frame.grid(row=2, column=1, sticky="ew", padx=(0, 8), pady=(10, 0))
        start_month_frame.columnconfigure(0, weight=1)
        self.cbo_start_month = ttk.Combobox(start_month_frame, textvariable=self.var_start_month, width=12)
        self.cbo_start_month.grid(row=0, column=0, sticky="ew")
        self.cbo_start_month.bind("<KeyRelease>", lambda _e: self.filter_combobox_values(self.cbo_start_month, self.month_options, self.var_start_month.get()))
        self.cbo_start_month.bind("<FocusOut>", lambda _e: self.var_start_month.set(self.normalize_month_text(self.var_start_month.get())))
        self.cbo_start_month.bind("<<ComboboxSelected>>", lambda _e: self.apply_filter(show_no_result_feedback=True))
        ttk.Button(start_month_frame, text="选", width=3, command=lambda: self.show_month_picker("start"), takefocus=0).grid(row=0, column=1, padx=(4, 0))

        ttk.Label(filter_frame, text="至", style="FormLabel.TLabel").grid(row=2, column=2, sticky="w", padx=(0, 8), pady=(10, 0))
        end_month_frame = ttk.Frame(filter_frame)
        end_month_frame.grid(row=2, column=3, sticky="ew", padx=(0, 18), pady=(10, 0))
        end_month_frame.columnconfigure(0, weight=1)
        self.cbo_end_month = ttk.Combobox(end_month_frame, textvariable=self.var_end_month, width=12)
        self.cbo_end_month.grid(row=0, column=0, sticky="ew")
        self.cbo_end_month.bind("<KeyRelease>", lambda _e: self.filter_combobox_values(self.cbo_end_month, self.month_options, self.var_end_month.get()))
        self.cbo_end_month.bind("<FocusOut>", lambda _e: self.var_end_month.set(self.normalize_month_text(self.var_end_month.get())))
        self.cbo_end_month.bind("<<ComboboxSelected>>", lambda _e: self.apply_filter(show_no_result_feedback=True))
        ttk.Button(end_month_frame, text="选", width=3, command=lambda: self.show_month_picker("end"), takefocus=0).grid(row=0, column=1, padx=(4, 0))

        ttk.Label(filter_frame, text="关键字", style="FormLabel.TLabel").grid(row=3, column=0, sticky="w", padx=(0, 8), pady=(10, 0))
        ttk.Entry(filter_frame, textvariable=self.var_filter_keyword).grid(row=3, column=1, columnspan=3, sticky="ew", padx=(0, 18), pady=(10, 0))
        ttk.Button(filter_frame, text="查询", command=self.apply_filter, style="Primary.TButton", takefocus=0).grid(row=3, column=4, padx=(0, 8), pady=(10, 0), sticky="ew")
        ttk.Button(filter_frame, text="清空", command=self.clear_filter, takefocus=0).grid(row=3, column=5, padx=(0, 0), pady=(10, 0), sticky="ew")
        list_frame = ttk.Frame(self.tab_invoice)
        list_frame.pack(fill=tk.BOTH, expand=True, padx=6, pady=(0, 8))

        columns = ("id", "seller", "buyer", "invoice_no", "amount", "date", "file", "imported_at")
        self.tree_invoice = ttk.Treeview(list_frame, columns=columns, show="headings", selectmode="extended")
        for key in columns:
            self.tree_invoice.heading(key, command=lambda current=key: self.sort_by_invoice_column(current))
        self.update_invoice_heading_labels()
        self.tree_invoice.column("id", width=60, anchor="center")
        self.tree_invoice.column("seller", width=200)
        self.tree_invoice.column("buyer", width=200)
        self.tree_invoice.column("invoice_no", width=170, anchor="center")
        self.tree_invoice.column("amount", width=110, anchor="e")
        self.tree_invoice.column("date", width=110, anchor="center")
        self.tree_invoice.column("file", width=320)
        self.tree_invoice.column("imported_at", width=150, anchor="center")
        self.tree_invoice.bind("<Double-1>", lambda _e: self.open_selected_invoice_file())
        self.tree_invoice.bind("<Configure>", lambda _e: self._sync_invoice_summary_layout())

        invoice_scroll = ttk.Scrollbar(list_frame, orient=tk.VERTICAL, command=self.tree_invoice.yview)
        self.invoice_xscroll = ttk.Scrollbar(list_frame, orient=tk.HORIZONTAL, command=self.tree_invoice.xview)
        self.tree_invoice.configure(yscrollcommand=invoice_scroll.set, xscrollcommand=self._on_invoice_xscroll)
        self.invoice_xscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree_invoice.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        invoice_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        summary_frame = ttk.Frame(self.tab_invoice, padding=(12, 8), style="White.TFrame")
        summary_frame.pack(fill=tk.X, padx=6, pady=(0, 8))
        summary_frame.configure(height=40)
        summary_frame.pack_propagate(False)
        self.summary_cells: dict[str, ttk.Label | ttk.Frame] = {}
        self.summary_cells["frame"] = summary_frame
        self.summary_cells["label"] = ttk.Label(summary_frame, text="当前筛选总计", style="SummaryLabel.TLabel", anchor="w")
        self.summary_cells["label"].place(x=12, rely=0.5, anchor="w")
        self.summary_cells["amount"] = ttk.Label(summary_frame, textvariable=self.var_total_amount, style="SummaryValue.TLabel", anchor="e")
        self.summary_cells["amount"].place(x=12, y=8, width=110, height=24)
        self._sync_invoice_summary_layout()

        action_frame = ttk.Frame(self.tab_invoice)
        action_frame.pack(fill=tk.X, padx=6, pady=(0, 6))
        ttk.Button(action_frame, text="批量修改字段", command=self.batch_edit_invoice_fields).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(action_frame, text="批量删除", command=self.delete_selected_invoices).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(action_frame, text="打开源文件", command=self.open_selected_invoice_file).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(action_frame, text="打开文件夹", command=self.open_selected_file_dir).pack(side=tk.LEFT)

        page_frame = ttk.Frame(self.tab_invoice)
        page_frame.pack(fill=tk.X, padx=6, pady=(0, 4))
        self.lbl_page = ttk.Label(page_frame, text="第 1 / 1 页")
        self.lbl_page.pack(side=tk.LEFT)
        ttk.Label(page_frame, text="每页").pack(side=tk.LEFT, padx=(18, 4))
        page_size_combo = ttk.Combobox(
            page_frame,
            textvariable=self.var_page_size,
            values=PAGE_SIZE_OPTIONS,
            width=6,
            state="readonly",
        )
        page_size_combo.pack(side=tk.LEFT)
        page_size_combo.bind("<<ComboboxSelected>>", lambda _e: self.on_page_size_change())

        nav_frame = ttk.Frame(page_frame)
        nav_frame.pack(side=tk.RIGHT)
        ttk.Button(nav_frame, text="上一页", command=self.prev_page).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(nav_frame, text="下一页", command=self.next_page).pack(side=tk.LEFT)

    def _build_approval_tab(self) -> None:
        top = ttk.Frame(self.tab_approval)
        top.pack(fill=tk.X, padx=6, pady=(8, 8))
        ttk.Button(top, text="刷新审批池", command=self.refresh_pending_table).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="批量审批", command=self.bulk_approve_selected).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="批量替换销售方", command=lambda: self.bulk_replace_pending_party("seller")).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="批量替换购买方", command=lambda: self.bulk_replace_pending_party("buyer")).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="批量删除", command=self.delete_selected_pending_records).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="清理失效文件", command=self.clean_missing_pending_records).pack(side=tk.LEFT, padx=(0, 6))



        paned = ttk.Panedwindow(self.tab_approval, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True, padx=6, pady=(0, 8))

        left = ttk.Frame(paned)
        right = ttk.Frame(paned, style="White.TFrame")
        paned.add(left, weight=3)
        paned.add(right, weight=2)

        columns = ("id", "seller", "buyer", "invoice_no", "amount", "date", "reason", "file")
        self.tree_pending = ttk.Treeview(left, columns=columns, show="headings", selectmode="extended")
        headings = {
            "id": "ID",
            "seller": "销售方",
            "buyer": "购买方",
            "invoice_no": "发票号码",
            "amount": "金额",
            "date": "开票日期",
            "reason": "审批原因",
            "file": "文件名",
        }
        for key, label in headings.items():
            self.tree_pending.heading(key, text=label)
        self.tree_pending.column("id", width=50, anchor="center")
        self.tree_pending.column("seller", width=180)
        self.tree_pending.column("buyer", width=180)
        self.tree_pending.column("invoice_no", width=160, anchor="center")
        self.tree_pending.column("amount", width=100, anchor="e")
        self.tree_pending.column("date", width=110, anchor="center")
        self.tree_pending.column("reason", width=220)
        self.tree_pending.column("file", width=260)
        self.tree_pending.bind("<<TreeviewSelect>>", lambda _e: self.on_pending_select())
        self.tree_pending.bind("<Double-1>", lambda _e: self.open_selected_pending_file())
        pending_scroll = ttk.Scrollbar(left, orient=tk.VERTICAL, command=self.tree_pending.yview)
        pending_xscroll = ttk.Scrollbar(left, orient=tk.HORIZONTAL, command=self.tree_pending.xview)
        self.tree_pending.configure(yscrollcommand=pending_scroll.set, xscrollcommand=pending_xscroll.set)
        pending_xscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree_pending.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        pending_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        detail_frame = ttk.LabelFrame(right, text="审批详情", padding=12, style="Card.TLabelframe")
        detail_frame.pack(fill=tk.BOTH, expand=True)
        detail_frame.columnconfigure(1, weight=1)

        ttk.Label(detail_frame, text="销售方").grid(row=0, column=0, sticky="w", pady=4)
        self.cbo_pending_seller = ttk.Combobox(detail_frame, textvariable=self.var_pending_seller)
        self.cbo_pending_seller.grid(row=0, column=1, sticky="ew", pady=4)
        self._bind_searchable_combobox(self.cbo_pending_seller, lambda: self.pending_seller_options, self.var_pending_seller)

        ttk.Label(detail_frame, text="购买方").grid(row=1, column=0, sticky="w", pady=4)
        self.cbo_pending_buyer = ttk.Combobox(detail_frame, textvariable=self.var_pending_buyer)
        self.cbo_pending_buyer.grid(row=1, column=1, sticky="ew", pady=4)
        self._bind_searchable_combobox(self.cbo_pending_buyer, lambda: self.pending_buyer_options, self.var_pending_buyer)

        ttk.Label(detail_frame, text="发票号").grid(row=2, column=0, sticky="w", pady=4)
        ttk.Entry(detail_frame, textvariable=self.var_pending_invoice_no, state="readonly").grid(row=2, column=1, sticky="ew", pady=4)
        ttk.Label(detail_frame, text="金额").grid(row=3, column=0, sticky="w", pady=4)
        ttk.Entry(detail_frame, textvariable=self.var_pending_amount, state="readonly").grid(row=3, column=1, sticky="ew", pady=4)
        ttk.Label(detail_frame, text="日期").grid(row=4, column=0, sticky="w", pady=4)
        ttk.Entry(detail_frame, textvariable=self.var_pending_date, state="readonly").grid(row=4, column=1, sticky="ew", pady=4)
        ttk.Label(detail_frame, text="审批原因").grid(row=5, column=0, sticky="nw", pady=4)
        ttk.Label(detail_frame, textvariable=self.var_pending_reason, wraplength=360, justify="left").grid(row=5, column=1, sticky="w", pady=4)
        ttk.Label(detail_frame, text="复核原因").grid(row=6, column=0, sticky="nw", pady=4)
        ttk.Label(detail_frame, textvariable=self.var_pending_review_reason, wraplength=360, justify="left").grid(row=6, column=1, sticky="w", pady=4)
        ttk.Label(detail_frame, text="文件").grid(row=7, column=0, sticky="nw", pady=4)
        ttk.Label(detail_frame, textvariable=self.var_pending_file, wraplength=360, justify="left").grid(row=7, column=1, sticky="w", pady=4)

        ttk.Label(detail_frame, text="销售方同步").grid(row=8, column=0, sticky="w", pady=(8, 0))
        seller_sync_opts = ttk.Frame(detail_frame)
        seller_sync_opts.grid(row=8, column=1, sticky="w", pady=(8, 0))
        ttk.Radiobutton(seller_sync_opts, text="不处理", variable=self.var_pending_seller_sync, value="none", command=lambda: self.on_pending_sync_mode_changed("seller")).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Radiobutton(seller_sync_opts, text="加入白名单", variable=self.var_pending_seller_sync, value="whitelist", command=lambda: self.on_pending_sync_mode_changed("seller")).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Radiobutton(seller_sync_opts, text="记为别名", variable=self.var_pending_seller_sync, value="alias", command=lambda: self.on_pending_sync_mode_changed("seller")).pack(side=tk.LEFT)

        ttk.Label(detail_frame, text="购买方同步").grid(row=9, column=0, sticky="w", pady=(6, 0))
        buyer_sync_opts = ttk.Frame(detail_frame)
        buyer_sync_opts.grid(row=9, column=1, sticky="w", pady=(6, 0))
        ttk.Radiobutton(buyer_sync_opts, text="不处理", variable=self.var_pending_buyer_sync, value="none", command=lambda: self.on_pending_sync_mode_changed("buyer")).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Radiobutton(buyer_sync_opts, text="加入白名单", variable=self.var_pending_buyer_sync, value="whitelist", command=lambda: self.on_pending_sync_mode_changed("buyer")).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Radiobutton(buyer_sync_opts, text="记为别名", variable=self.var_pending_buyer_sync, value="alias", command=lambda: self.on_pending_sync_mode_changed("buyer")).pack(side=tk.LEFT)

        ttk.Label(detail_frame, text="原文参考").grid(row=10, column=0, sticky="nw", pady=(12, 4))
        self.txt_pending_raw = tk.Text(
            detail_frame,
            height=18,
            wrap="word",
            font=(self.font_family, 9),
            bg="#FFFFFF",
            fg="#1D1D1F",
            relief="flat",
            highlightthickness=1,
            highlightbackground="#E6EAF0",
            highlightcolor="#D8DEE8",
        )
        self.txt_pending_raw.grid(row=10, column=1, sticky="nsew", pady=(12, 4))
        detail_frame.rowconfigure(10, weight=1)

        btn_frame = ttk.Frame(detail_frame)
        btn_frame.grid(row=11, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        ttk.Button(btn_frame, text="确认", command=self.confirm_current_pending, style="Primary.TButton").pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(btn_frame, text="删除选中", command=self.delete_selected_pending_records).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(btn_frame, text="打开源文件", command=self.open_selected_pending_file).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(btn_frame, text="打开文件夹", command=self.open_selected_pending_folder).pack(side=tk.LEFT)

    def _build_whitelist_tab(self) -> None:
        notebook = ttk.Notebook(self.tab_whitelist)
        notebook.pack(fill=tk.BOTH, expand=True, padx=6, pady=8)
        seller_frame = ttk.Frame(notebook)
        buyer_frame = ttk.Frame(notebook)
        notebook.add(seller_frame, text="销售方白名单")
        notebook.add(buyer_frame, text="购买方白名单")
        self._build_whitelist_panel("seller", seller_frame)
        self._build_whitelist_panel("buyer", buyer_frame)

    def _build_whitelist_panel(self, party_type: str, parent: ttk.Frame) -> None:
        top = ttk.Frame(parent)
        top.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(top, text="标准名搜索").pack(side=tk.LEFT, padx=(0, 6))
        entry_search = ttk.Entry(top, textvariable=self.whitelist_search_vars[party_type], width=24)
        entry_search.pack(side=tk.LEFT, padx=(0, 12))
        entry_search.bind("<KeyRelease>", lambda _e, p=party_type: self.refresh_whitelist_view(p))
        ttk.Label(top, text="别名搜索").pack(side=tk.LEFT, padx=(0, 6))
        alias_search = ttk.Entry(top, textvariable=self.whitelist_alias_search_vars[party_type], width=20)
        alias_search.pack(side=tk.LEFT, padx=(0, 12))
        alias_search.bind("<KeyRelease>", lambda _e, p=party_type: self.refresh_whitelist_alias_view(p))
        ttk.Button(top, text="新增标准名", command=lambda: self.add_whitelist_entry(party_type), style="Primary.TButton").pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="编辑标准名", command=lambda: self.edit_whitelist_entry(party_type)).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="删除标准名", command=lambda: self.delete_whitelist_entry(party_type)).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="导入Excel", command=lambda: self.import_whitelist_excel(party_type)).pack(side=tk.LEFT, padx=(0, 12))
        ttk.Button(top, text="新增别名", command=lambda: self.add_whitelist_alias(party_type)).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="编辑别名", command=lambda: self.edit_whitelist_alias(party_type)).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(top, text="删除别名", command=lambda: self.delete_whitelist_alias(party_type)).pack(side=tk.LEFT)

        paned = ttk.Panedwindow(parent, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True)
        left = ttk.Frame(paned)
        right = ttk.Frame(paned)
        paned.add(left, weight=3)
        paned.add(right, weight=2)

        tree_entries = ttk.Treeview(left, columns=("id", "name"), show="headings", selectmode="browse")
        tree_entries.heading("id", text="ID", command=lambda p=party_type: self.sort_whitelist_entries(p, "id"))
        tree_entries.heading("name", text="标准名", command=lambda p=party_type: self.sort_whitelist_entries(p, "name"))
        tree_entries.column("id", width=60, anchor="center")
        tree_entries.column("name", width=360)
        tree_entries.bind("<<TreeviewSelect>>", lambda _e, p=party_type: self.on_whitelist_select(p))
        scroll_entries = ttk.Scrollbar(left, orient=tk.VERTICAL, command=tree_entries.yview)
        xscroll_entries = ttk.Scrollbar(left, orient=tk.HORIZONTAL, command=tree_entries.xview)
        tree_entries.configure(yscrollcommand=scroll_entries.set, xscrollcommand=xscroll_entries.set)
        xscroll_entries.pack(side=tk.BOTTOM, fill=tk.X)
        tree_entries.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll_entries.pack(side=tk.RIGHT, fill=tk.Y)

        tree_aliases = ttk.Treeview(right, columns=("id", "alias"), show="headings", selectmode="browse")
        tree_aliases.heading("id", text="ID", command=lambda p=party_type: self.sort_whitelist_aliases(p, "id"))
        tree_aliases.heading("alias", text="别名", command=lambda p=party_type: self.sort_whitelist_aliases(p, "alias"))
        tree_aliases.column("id", width=60, anchor="center")
        tree_aliases.column("alias", width=300)
        scroll_aliases = ttk.Scrollbar(right, orient=tk.VERTICAL, command=tree_aliases.yview)
        xscroll_aliases = ttk.Scrollbar(right, orient=tk.HORIZONTAL, command=tree_aliases.xview)
        tree_aliases.configure(yscrollcommand=scroll_aliases.set, xscrollcommand=xscroll_aliases.set)
        xscroll_aliases.pack(side=tk.BOTTOM, fill=tk.X)
        tree_aliases.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll_aliases.pack(side=tk.RIGHT, fill=tk.Y)

        self.whitelist_widgets[party_type] = {
            "entries": tree_entries,
            "aliases": tree_aliases,
            "entry_sort_key": "name",
            "entry_sort_ascending": True,
            "alias_sort_key": "alias",
            "alias_sort_ascending": True,
        }

    def _build_reconcile_tab(self) -> None:
        notebook = ttk.Notebook(self.tab_reconcile)
        notebook.pack(fill=tk.BOTH, expand=True, padx=6, pady=8)

        tab_manual = ttk.Frame(notebook)
        tab_batch = ttk.Frame(notebook)
        notebook.add(tab_manual, text="手动核对")
        notebook.add(tab_batch, text="批量对账")

        manual_form = ttk.LabelFrame(tab_manual, text="核对条件", padding=10, style="Card.TLabelframe")
        manual_form.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(manual_form, text="销售方").grid(row=0, column=0, sticky="w")
        self.cbo_rec_seller = ttk.Combobox(manual_form, textvariable=self.var_rec_seller, width=24)
        self.cbo_rec_seller.grid(row=0, column=1, sticky="w", padx=(4, 12))
        self.cbo_rec_seller.bind("<KeyRelease>", lambda _e: self.filter_combobox_values(self.cbo_rec_seller, self.reconcile_seller_options, self.var_rec_seller.get()))
        ttk.Label(manual_form, text="购买方").grid(row=0, column=2, sticky="w")
        self.cbo_rec_buyer = ttk.Combobox(manual_form, textvariable=self.var_rec_buyer, width=24)
        self.cbo_rec_buyer.grid(row=0, column=3, sticky="w", padx=(4, 12))
        self.cbo_rec_buyer.bind("<KeyRelease>", lambda _e: self.filter_combobox_values(self.cbo_rec_buyer, self.reconcile_buyer_options, self.var_rec_buyer.get()))
        ttk.Label(manual_form, text="发票号").grid(row=0, column=4, sticky="w")
        ttk.Entry(manual_form, textvariable=self.var_rec_no, width=22).grid(row=0, column=5, sticky="w", padx=(4, 12))
        ttk.Label(manual_form, text="金额").grid(row=0, column=6, sticky="w")
        ttk.Entry(manual_form, textvariable=self.var_rec_amount, width=16).grid(row=0, column=7, sticky="w", padx=(4, 12))
        ttk.Button(manual_form, text="执行核对", command=self.run_reconcile, style="Primary.TButton").grid(row=0, column=8)

        manual_tip = ttk.Label(
            tab_manual,
            text="适合临时核对单张或少量发票。按销售方/购买方/发票号/金额查询正式库，并标出是否匹配。",
            anchor="w",
        )
        manual_tip.pack(fill=tk.X, pady=(0, 8))

        result_frame = ttk.Frame(tab_manual)
        result_frame.pack(fill=tk.BOTH, expand=True)
        columns = ("seller", "buyer", "invoice_no", "amount", "date", "status", "file")
        self.tree_reconcile = ttk.Treeview(result_frame, columns=columns, show="headings")
        for key in columns:
            self.tree_reconcile.heading(key, command=lambda current=key: self.sort_reconcile_column(current))
        self.update_reconcile_heading_labels()
        self.tree_reconcile.column("seller", width=180)
        self.tree_reconcile.column("buyer", width=180)
        self.tree_reconcile.column("invoice_no", width=160, anchor="center")
        self.tree_reconcile.column("amount", width=100, anchor="e")
        self.tree_reconcile.column("date", width=110, anchor="center")
        self.tree_reconcile.column("status", width=140, anchor="center")
        self.tree_reconcile.column("file", width=320)
        self.tree_reconcile.bind("<Double-1>", lambda _e: self.open_selected_reconcile_file())
        scroll = ttk.Scrollbar(result_frame, orient=tk.VERTICAL, command=self.tree_reconcile.yview)
        xscroll = ttk.Scrollbar(result_frame, orient=tk.HORIZONTAL, command=self.tree_reconcile.xview)
        self.tree_reconcile.configure(yscrollcommand=scroll.set, xscrollcommand=xscroll.set)
        xscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree_reconcile.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scroll.pack(side=tk.RIGHT, fill=tk.Y)

        batch_top = ttk.LabelFrame(tab_batch, text="批量对账", padding=10, style="Card.TLabelframe")
        batch_top.pack(fill=tk.X, pady=(0, 8))
        ttk.Label(
            batch_top,
            text="用于导入供应商发来的发票清单。模板建议至少包含“发票号”，可选填写销售方、购买方、金额。",
            anchor="w",
        ).pack(fill=tk.X, pady=(0, 8))
        btns = ttk.Frame(batch_top)
        btns.pack(fill=tk.X)
        ttk.Button(btns, text="下载模板", command=self.download_reconcile_template).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(btns, text="导入对账清单", command=self.import_reconcile_sheet, style="Primary.TButton").pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(btns, text="导出对账结果", command=self.export_reconcile_results).pack(side=tk.LEFT)
        ttk.Checkbutton(btns, text="只看异常项", variable=self.var_batch_reconcile_only_issues, command=self._populate_batch_reconcile_tree).pack(side=tk.LEFT, padx=(16, 0))

        batch_frame = ttk.Frame(tab_batch)
        batch_frame.pack(fill=tk.BOTH, expand=True)
        batch_columns = (
            "input_seller",
            "input_buyer",
            "input_invoice_no",
            "input_amount",
            "result",
            "matched_seller",
            "matched_buyer",
            "matched_amount",
            "matched_date",
            "file",
        )
        self.tree_batch_reconcile = ttk.Treeview(batch_frame, columns=batch_columns, show="headings")
        for key in batch_columns:
            self.tree_batch_reconcile.heading(key, command=lambda current=key: self.sort_batch_reconcile_column(current))
        self.update_batch_reconcile_heading_labels()
        self.tree_batch_reconcile.column("input_seller", width=170)
        self.tree_batch_reconcile.column("input_buyer", width=170)
        self.tree_batch_reconcile.column("input_invoice_no", width=150, anchor="center")
        self.tree_batch_reconcile.column("input_amount", width=100, anchor="e")
        self.tree_batch_reconcile.column("result", width=140, anchor="center")
        self.tree_batch_reconcile.column("matched_seller", width=170)
        self.tree_batch_reconcile.column("matched_buyer", width=170)
        self.tree_batch_reconcile.column("matched_amount", width=100, anchor="e")
        self.tree_batch_reconcile.column("matched_date", width=110, anchor="center")
        self.tree_batch_reconcile.column("file", width=260)
        self.tree_batch_reconcile.bind("<Double-1>", lambda _e: self.open_selected_batch_reconcile_file())
        batch_scroll = ttk.Scrollbar(batch_frame, orient=tk.VERTICAL, command=self.tree_batch_reconcile.yview)
        batch_xscroll = ttk.Scrollbar(batch_frame, orient=tk.HORIZONTAL, command=self.tree_batch_reconcile.xview)
        self.tree_batch_reconcile.configure(yscrollcommand=batch_scroll.set, xscrollcommand=batch_xscroll.set)
        batch_xscroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.tree_batch_reconcile.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        batch_scroll.pack(side=tk.RIGHT, fill=tk.Y)

    def set_status(self, msg: str) -> None:
        self.var_status.set(msg)
        self.update_idletasks()

    def refresh_all(self) -> None:
        self.refresh_filter_options()
        self.refresh_invoice_table()
        self.refresh_pending_table()
        self.refresh_whitelist_views()

    def refresh_filter_options(self) -> None:
        self.invoice_seller_options = self.store.get_all_invoice_party_names("seller", pool="approved")
        self.invoice_buyer_options = self.store.get_all_invoice_party_names("buyer", pool="approved")
        self.month_options = self.store.get_all_months(pool="approved")
        self.pending_seller_options = self.store.get_all_party_names("seller")
        self.pending_buyer_options = self.store.get_all_party_names("buyer")
        self.reconcile_seller_options = sorted(set(self.invoice_seller_options) | set(self.store.get_all_party_names("seller")))
        self.reconcile_buyer_options = sorted(set(self.invoice_buyer_options) | set(self.store.get_all_party_names("buyer")))

        self.cbo_filter_seller["values"] = self.invoice_seller_options
        self.cbo_filter_buyer["values"] = self.invoice_buyer_options
        self.cbo_start_month["values"] = self.month_options
        self.cbo_end_month["values"] = self.month_options
        self.cbo_pending_seller["values"] = self.pending_seller_options
        self.cbo_pending_buyer["values"] = self.pending_buyer_options
        self.cbo_rec_seller["values"] = self.reconcile_seller_options
        self.cbo_rec_buyer["values"] = self.reconcile_buyer_options


    def filter_combobox_values(self, combo: ttk.Combobox, source_values: list[str], current_text: str) -> None:
        keyword = clean_company_name(current_text)
        if not keyword:
            combo["values"] = source_values
            try:
                combo.tk.call("ttk::combobox::Unpost", str(combo))
            except tk.TclError:
                pass
            return
        compact = keyword.upper()
        matched = [item for item in source_values if compact in clean_company_name(item).upper()]
        combo["values"] = matched
        if matched:
            self.after_idle(lambda: self._show_combobox_dropdown(combo))
        else:
            try:
                combo.tk.call("ttk::combobox::Unpost", str(combo))
            except tk.TclError:
                pass

    @staticmethod
    def _show_combobox_dropdown(combo: ttk.Combobox) -> None:
        try:
            combo.focus_set()
            combo.tk.call("ttk::combobox::Post", str(combo))
        except tk.TclError:
            pass

    def _commit_combobox_selection(self, combo: ttk.Combobox, var: tk.StringVar | None = None) -> None:
        value = clean_company_name(combo.get())
        if var is not None:
            var.set(value)
        else:
            combo.set(value)
        try:
            combo.tk.call("ttk::combobox::Unpost", str(combo))
        except tk.TclError:
            pass
        self.after_idle(self.focus_set)

    def _bind_searchable_combobox(
        self,
        combo: ttk.Combobox,
        source_getter,
        var: tk.StringVar | None = None,
    ) -> None:
        combo.bind(
            "<KeyRelease>",
            lambda _e: self.filter_combobox_values(combo, list(source_getter()), combo.get()),
        )
        combo.bind("<<ComboboxSelected>>", lambda _e: self._commit_combobox_selection(combo, var))
        combo.bind("<Return>", lambda _e: self._commit_combobox_selection(combo, var))

    def on_pending_sync_mode_changed(self, party_type: str) -> None:
        if party_type == "seller":
            combo = self.cbo_pending_seller
            mode = self.var_pending_seller_sync.get()
            options = self.pending_seller_options
            label = "销售方"
        else:
            combo = self.cbo_pending_buyer
            mode = self.var_pending_buyer_sync.get()
            options = self.pending_buyer_options
            label = "购买方"

        combo["values"] = options
        if mode == "alias":
            if not options:
                messagebox.showwarning("提示", f"当前没有{label}白名单，不能记为别名。请先新增标准白名单。")
                if party_type == "seller":
                    self.var_pending_seller_sync.set("none")
                else:
                    self.var_pending_buyer_sync.set("none")
                return
            self.set_status(f"记为别名：请在{label}输入框选择已有标准白名单名称")
            self.after_idle(lambda: self._show_combobox_dropdown(combo))

    def get_page_size(self) -> int:
        try:
            value = int(self.var_page_size.get())
        except (TypeError, ValueError, tk.TclError):
            return DEFAULT_PAGE_SIZE
        return value if value in PAGE_SIZE_OPTIONS else DEFAULT_PAGE_SIZE

    def on_page_size_change(self) -> None:
        self.current_page = 1
        self.refresh_invoice_table(show_no_result_feedback=False)

    def _on_invoice_xscroll(self, first: str, last: str) -> None:
        if hasattr(self, "invoice_xscroll"):
            self.invoice_xscroll.set(first, last)
        self._sync_invoice_summary_layout()

    def get_invoice_filters(self) -> tuple[str, str, str, str, str]:
        return (
            self.var_filter_seller.get().strip(),
            self.var_filter_buyer.get().strip(),
            self.normalize_month_text(self.var_start_month.get().strip()),
            self.normalize_month_text(self.var_end_month.get().strip()),
            self.var_filter_keyword.get().strip(),
        )


    def update_invoice_heading_labels(self) -> None:
        for key, label in self.invoice_heading_labels.items():
            heading = label
            if key == self.sort_column:
                heading = f"{label} {'▲' if self.sort_ascending else '▼'}"
            self.tree_invoice.heading(key, text=heading)

    def update_reconcile_heading_labels(self) -> None:
        for key, label in self.reconcile_heading_labels.items():
            heading = label
            if key == self.reconcile_sort_column:
                heading = f"{label} {'▲' if self.reconcile_sort_ascending else '▼'}"
            self.tree_reconcile.heading(key, text=heading)

    def update_batch_reconcile_heading_labels(self) -> None:
        for key, label in self.batch_reconcile_heading_labels.items():
            heading = label
            if key == self.batch_reconcile_sort_column:
                heading = f"{label} {'▲' if self.batch_reconcile_sort_ascending else '▼'}"
            self.tree_batch_reconcile.heading(key, text=heading)

    def sort_by_invoice_column(self, column_key: str) -> None:
        if self.sort_column == column_key:
            self.sort_ascending = not self.sort_ascending
        else:
            self.sort_column = column_key
            self.sort_ascending = True
        self.current_page = 1
        self.update_invoice_heading_labels()
        self.refresh_invoice_table()

    def sort_reconcile_column(self, column_key: str) -> None:
        if self.reconcile_sort_column == column_key:
            self.reconcile_sort_ascending = not self.reconcile_sort_ascending
        else:
            self.reconcile_sort_column = column_key
            self.reconcile_sort_ascending = True
        self.update_reconcile_heading_labels()
        self._populate_reconcile_tree()

    def sort_batch_reconcile_column(self, column_key: str) -> None:
        if self.batch_reconcile_sort_column == column_key:
            self.batch_reconcile_sort_ascending = not self.batch_reconcile_sort_ascending
        else:
            self.batch_reconcile_sort_column = column_key
            self.batch_reconcile_sort_ascending = True
        self.update_batch_reconcile_heading_labels()
        self._populate_batch_reconcile_tree()

    def apply_filter(self, show_no_result_feedback: bool = True) -> None:
        start_month = self.normalize_month_text(self.var_start_month.get().strip())
        end_month = self.normalize_month_text(self.var_end_month.get().strip())
        self.var_start_month.set(start_month)
        self.var_end_month.set(end_month)
        if start_month and not self._is_valid_month(start_month):
            messagebox.showwarning("格式提示", "开始月份请使用 YYYY-MM，例如 2026-04")
            return
        if end_month and not self._is_valid_month(end_month):
            messagebox.showwarning("格式提示", "结束月份请使用 YYYY-MM，例如 2026-04")
            return
        if start_month and end_month and start_month > end_month:
            messagebox.showwarning("格式提示", "开始月份不能晚于结束月份")
            return
        self.current_page = 1
        self.refresh_invoice_table(show_no_result_feedback=show_no_result_feedback)

    def clear_filter(self) -> None:
        self.var_filter_seller.set("")
        self.var_filter_buyer.set("")
        self.var_start_month.set("")
        self.var_end_month.set("")

        self.var_filter_keyword.set("")
        self.current_page = 1
        self.refresh_invoice_table(show_no_result_feedback=False)

    def refresh_invoice_table(self, show_no_result_feedback: bool = False) -> None:
        seller, buyer, start_month, end_month, keyword = self.get_invoice_filters()
        page_size = self.get_page_size()
        rows, total = self.store.query_invoices(
            seller_keyword=seller,
            buyer_keyword=buyer,
            start_month=start_month,
            end_month=end_month,
            keyword=keyword,
            page=self.current_page,
            page_size=page_size,
            order_by=self.sort_column,
            ascending=self.sort_ascending,
            pool="approved",
        )
        total_amount = self.store.get_filtered_total_amount(
            seller,
            buyer,
            start_month,
            end_month,
            "",
            keyword,
            pool="approved",
        )
        self.total_items = total
        total_pages = max(1, (total + page_size - 1) // page_size)
        if total > 0 and self.current_page > total_pages:
            self.current_page = total_pages
            rows, total = self.store.query_invoices(
                seller_keyword=seller,
                buyer_keyword=buyer,
                start_month=start_month,
                end_month=end_month,
                keyword=keyword,
                page=self.current_page,
                page_size=page_size,
                order_by=self.sort_column,
                ascending=self.sort_ascending,
                pool="approved",
            )
            self.total_items = total

        for item_id in self.tree_invoice.get_children():
            self.tree_invoice.delete(item_id)
        for item in rows:
            self.tree_invoice.insert(
                "",
                tk.END,
                iid=str(item.id),
                values=(
                    item.id,
                    item.seller_name,
                    item.buyer_name,
                    item.invoice_number,
                    f"{item.amount:.2f}",
                    item.invoice_date,
                    item.stored_filename,
                    item.imported_at,
                ),
            )

        self.lbl_page.configure(text=f"第 {self.current_page} / {total_pages} 页  共 {total} 条")
        self.var_total_amount.set(f"¥ {total_amount:,.2f}")
        self._sync_invoice_summary_layout()
        filters_active = any((seller, buyer, start_month, end_month, keyword))
        if total == 0 and filters_active:
            self.set_status("当前条件下没有匹配记录")
            if show_no_result_feedback:
                messagebox.showinfo("查询结果", "当前条件下没有匹配记录")
        else:
            self.set_status(f"已加载 {len(rows)} 条记录")





    def refresh_pending_table(self) -> None:
        rows = self.store.get_pending_invoices()
        for item_id in self.tree_pending.get_children():
            self.tree_pending.delete(item_id)
        for item in rows:
            resolved = self._resolve_invoice_path(item, pool="pending", update_store=True)
            file_name = item.stored_filename if resolved else f"[文件缺失] {item.stored_filename}"
            self.tree_pending.insert(
                "",
                tk.END,
                iid=str(item.id),
                values=(
                    item.id,
                    item.seller_name,
                    item.buyer_name,
                    item.invoice_number,
                    f"{item.amount:.2f}",
                    item.invoice_date,
                    item.approval_reason,
                    file_name,
                ),
            )
        self.on_pending_select()

    def on_pending_select(self) -> None:
        selected = self.tree_pending.selection()
        if not selected:
            self.var_pending_seller.set("")
            self.var_pending_buyer.set("")
            self.var_pending_invoice_no.set("")
            self.var_pending_amount.set("")
            self.var_pending_date.set("")
            self.var_pending_reason.set("")
            self.var_pending_review_reason.set("")
            self.var_pending_file.set("")
            self.var_pending_seller_sync.set("none")
            self.var_pending_buyer_sync.set("none")
            self.txt_pending_raw.delete("1.0", tk.END)
            return
        item = self.store.get_invoice(int(selected[0]), pool="pending")
        if not item:
            return
        self.var_pending_seller.set(item.seller_name)
        self.var_pending_buyer.set(item.buyer_name)
        self.var_pending_invoice_no.set(item.invoice_number)
        self.var_pending_amount.set(f"{item.amount:.2f}")
        self.var_pending_date.set(item.invoice_date)
        self.var_pending_reason.set(item.approval_reason or "-")
        self.var_pending_review_reason.set(item.review_reason or "-")
        self.var_pending_file.set(item.stored_filename)
        self.var_pending_seller_sync.set("none")
        self.var_pending_buyer_sync.set("none")
        self.txt_pending_raw.delete("1.0", tk.END)
        self.txt_pending_raw.insert("1.0", item.raw_text or "(无原文文本)")

    def prev_page(self) -> None:
        if self.current_page > 1:
            self.current_page -= 1
            self.refresh_invoice_table()

    def next_page(self) -> None:
        page_size = self.get_page_size()
        total_pages = max(1, (self.total_items + page_size - 1) // page_size)
        if self.current_page < total_pages:
            self.current_page += 1
            self.refresh_invoice_table()

    def import_files(self) -> None:
        file_paths = filedialog.askopenfilenames(title="选择发票PDF", filetypes=[("PDF files", "*.pdf")])
        if file_paths:
            self._import_pdf_paths([Path(path) for path in file_paths])

    def import_folder(self) -> None:
        folder = filedialog.askdirectory(title="选择包含发票的文件夹")
        if not folder:
            return
        pdfs = [path for path in Path(folder).rglob("*.pdf")]
        if not pdfs:
            messagebox.showinfo("提示", "该文件夹下没有 PDF 文件")
            return
        self._import_pdf_paths(pdfs)

    def _import_pdf_paths(self, paths: list[Path]) -> None:
        seller_map = self.store.get_whitelist_map("seller")
        buyer_map = self.store.get_whitelist_map("buyer")
        batch_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        report_rows: list[dict[str, str]] = []

        for index, pdf_path in enumerate(paths, start=1):
            self.set_status(f"正在导入 {index}/{len(paths)}: {pdf_path.name}")
            try:
                parsed = parse_invoice(pdf_path, seller_map, buyer_map)
                seller_name = parsed.seller_name
                buyer_name = parsed.buyer_name
                invoice_no = parsed.invoice_number
                amount = float(parsed.amount or 0)
                invoice_date = parsed.invoice_date

                approval_reasons = [
                    reason
                    for reason in parsed.issues
                    if reason in {
                        "新销售方",
                        "新购买方",
                        "销售方无效",
                        "购买方无效",
                        "发票号码缺失",
                        "开票日期缺失",
                        "销售方与购买方相同",
                        "版式不匹配",
                        "主体定位低置信度",
                        "购买方区域未识别",
                        "销售方区域未识别",
                        "购买方区域未命中白名单",
                        "销售方区域未命中白名单",
                        "OCR解析异常",
                    }
                ]
                review_reasons = [reason for reason in parsed.issues if reason not in approval_reasons]
                if amount > AMOUNT_REVIEW_THRESHOLD:
                    review_reasons.append("金额大于10万")
                if parsed.ocr_used:
                    review_reasons.append("已触发OCR增强")

                if seller_name and buyer_name and invoice_no and self.store.invoice_exists(seller_name, buyer_name, invoice_no, pool="all"):
                    report_rows.append(
                        self._build_report_row("重复跳过", pdf_path.name, seller_name, buyer_name, invoice_no, amount, invoice_date, "重复发票")
                    )
                    continue

                approval_status = "normal" if not approval_reasons and parsed.seller_in_whitelist and parsed.buyer_in_whitelist else "pending"
                approval_reason = "; ".join(dict.fromkeys(approval_reasons))
                review_reason = "; ".join(dict.fromkeys(review_reasons))
                needs_review = 1 if review_reason else 0

                if approval_status == "normal":
                    stored_name = self._make_final_filename(seller_name, buyer_name, invoice_no)
                    target_path = self._copy_into_invoice_dir(pdf_path, stored_name, self.approved_invoice_dir)
                    file_finalized = 1
                    pool = "approved"
                else:
                    target_path = self._copy_into_invoice_dir(pdf_path, pdf_path.name, self.pending_invoice_dir)
                    stored_name = target_path.name
                    file_finalized = 0
                    pool = "pending"

                self.store.add_invoice(
                    seller_name=seller_name,
                    buyer_name=buyer_name,
                    invoice_number=invoice_no,
                    amount=amount,
                    invoice_date=invoice_date,
                    source_path=str(target_path),
                    original_filename=pdf_path.name,
                    stored_filename=stored_name,
                    needs_review=needs_review,
                    review_reason=review_reason,
                    approval_status=approval_status,
                    approval_reason=approval_reason,
                    batch_id=batch_id,
                    file_finalized=file_finalized,
                    raw_text=parsed.raw_text,
                    pool=pool,
                )
                report_rows.append(
                    self._build_report_row(
                        "正常入库" if approval_status == "normal" else "待审批",
                        pdf_path.name,
                        seller_name,
                        buyer_name,
                        invoice_no,
                        amount,
                        invoice_date,
                        approval_reason or review_reason or "正常",
                    )
                )
            except Exception as exc:
                report_rows.append(self._build_report_row("失败", pdf_path.name, "", "", "", 0, "", str(exc)))

        self.current_page = 1
        self.refresh_all()
        self.show_import_report(batch_id, report_rows)

    def _build_report_row(
        self,
        status: str,
        filename: str,
        seller_name: str,
        buyer_name: str,
        invoice_no: str,
        amount: float,
        invoice_date: str,
        reason: str,
    ) -> dict[str, str]:
        return {
            "status": status,
            "filename": filename,
            "seller_name": seller_name,
            "buyer_name": buyer_name,
            "invoice_no": invoice_no,
            "amount": f"{amount:.2f}" if amount else "",
            "invoice_date": invoice_date,
            "reason": reason,
        }

    def show_import_report(self, batch_id: str, rows: list[dict[str, str]]) -> None:
        dialog = tk.Toplevel(self)
        dialog.title(f"导入结果 - 批次 {batch_id}")
        dialog.geometry("1120x560")
        dialog.transient(self)

        columns = ("status", "filename", "seller_name", "buyer_name", "invoice_no", "amount", "invoice_date", "reason")
        tree = ttk.Treeview(dialog, columns=columns, show="headings")
        labels = {
            "status": "结果",
            "filename": "原文件名",
            "seller_name": "销售方",
            "buyer_name": "购买方",
            "invoice_no": "发票号",
            "amount": "金额",
            "invoice_date": "日期",
            "reason": "原因",
        }
        for key, label in labels.items():
            tree.heading(key, text=label)
        tree.column("status", width=90, anchor="center")
        tree.column("filename", width=240)
        tree.column("seller_name", width=180)
        tree.column("buyer_name", width=180)
        tree.column("invoice_no", width=150, anchor="center")
        tree.column("amount", width=90, anchor="e")
        tree.column("invoice_date", width=100, anchor="center")
        tree.column("reason", width=280)
        for row in rows:
            tree.insert("", tk.END, values=tuple(row[key] for key in columns))
        scroll = ttk.Scrollbar(dialog, orient=tk.VERTICAL, command=tree.yview)
        tree.configure(yscrollcommand=scroll.set)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(10, 0), pady=10)
        scroll.pack(side=tk.RIGHT, fill=tk.Y, pady=10, padx=(0, 10))

        def copy_report() -> None:
            lines = ["\t".join(labels[key] for key in columns)]
            lines.extend("\t".join(row[key] for key in columns) for row in rows)
            self.clipboard_clear()
            self.clipboard_append("\n".join(lines))
            self.set_status("导入结果已复制到剪贴板")

        def export_report() -> None:
            target = filedialog.asksaveasfilename(
                title="导出导入结果",
                initialfile=f"导入结果_{batch_id}.csv",
                defaultextension=".csv",
                filetypes=[("CSV", "*.csv")],
            )
            if not target:
                return
            with Path(target).open("w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow([labels[key] for key in columns])
                for row in rows:
                    writer.writerow([row[key] for key in columns])
            self.set_status(f"导入结果已导出到 {target}")

        btns = ttk.Frame(dialog)
        btns.pack(fill=tk.X, padx=10, pady=(0, 10))
        ttk.Button(btns, text="复制清单", command=copy_report).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(btns, text="导出CSV", command=export_report).pack(side=tk.LEFT)

    def _copy_into_invoice_dir(self, src: Path, preferred_name: str, target_dir: Path) -> Path:
        target_dir.mkdir(parents=True, exist_ok=True)
        preferred_path = target_dir / sanitize_filename(preferred_name)
        target_path = self._next_available_file(preferred_path)
        shutil.copy2(src, target_path)
        return target_path

    def _make_final_filename(self, seller_name: str, buyer_name: str, invoice_no: str) -> str:
        return sanitize_filename(f"{seller_name}_{buyer_name}_{invoice_no}.pdf")

    def _next_available_file(self, path: Path) -> Path:
        if not path.exists():
            return path
        stem = path.stem
        suffix = path.suffix
        for index in range(2, 10000):
            candidate = path.with_name(f"{stem}_{index}{suffix}")
            if not candidate.exists():
                return candidate
        return path.with_name(f"{stem}_{datetime.now().strftime('%H%M%S')}{suffix}")

    def _get_selected_invoice_ids(self) -> list[int]:
        return [int(item_id) for item_id in self.tree_invoice.selection()]

    def _get_selected_pending_ids(self) -> list[int]:
        return [int(item_id) for item_id in self.tree_pending.selection()]


    def _delete_pending_records(self, ids: list[int], remove_files: bool = True) -> tuple[int, list[str]]:
        rows = self.store.get_invoices(ids, pool="pending")
        if not rows:
            return 0, []
        deleted_ids: list[int] = []
        errors: list[str] = []
        for row in rows:
            resolved = self._resolve_invoice_path(row, pool="pending", update_store=False)
            if remove_files and resolved and resolved.exists():
                try:
                    resolved.unlink()
                except OSError as exc:
                    errors.append(f"{row.stored_filename}: {exc}")
                    continue
            deleted_ids.append(row.id)
        if deleted_ids:
            self.store.delete_invoices(deleted_ids, pool="pending")
        return len(deleted_ids), errors

    def _resolve_invoice_path(self, item: InvoiceRecord, pool: str, update_store: bool = True) -> Path | None:
        current_path = Path(item.source_path) if item.source_path else None
        if current_path and current_path.exists():
            return current_path

        primary_dir = self.approved_invoice_dir if pool == "approved" else self.pending_invoice_dir
        candidates = [
            primary_dir / item.stored_filename,
            self.invoice_root_dir / item.stored_filename,
            self.data_root / "invoice_files" / pool / item.stored_filename,
            self.runtime_root / "_internal" / "invoice_files" / item.stored_filename,
        ]
        for candidate in candidates:
            if candidate.exists():
                if update_store:
                    self.store.update_invoice_file(item.id, str(candidate), candidate.name, item.file_finalized, pool=pool)
                return candidate
        return None

    def _open_path(self, path: Path) -> None:
        if not path.exists():
            messagebox.showwarning("提示", f"文件不存在：{path}")
            return
        try:
            os.startfile(str(path))
        except AttributeError:
            subprocess.Popen(["open", str(path)])

    def open_selected_invoice_file(self) -> None:
        ids = self._get_selected_invoice_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择发票记录")
            return
        item = self.store.get_invoice(ids[0], pool="approved")
        if item:
            resolved = self._resolve_invoice_path(item, pool="approved")
            if resolved:
                self._open_path(resolved)
            else:
                messagebox.showwarning("提示", f"文件不存在：{item.stored_filename}")

    def open_selected_file_dir(self) -> None:
        ids = self._get_selected_invoice_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择发票记录")
            return
        item = self.store.get_invoice(ids[0], pool="approved")
        if item:
            resolved = self._resolve_invoice_path(item, pool="approved")
            folder = resolved.parent if resolved else self.approved_invoice_dir
            subprocess.Popen(["explorer", str(folder)])

    def open_selected_pending_file(self) -> None:
        ids = self._get_selected_pending_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择待审批记录")
            return
        item = self.store.get_invoice(ids[0], pool="pending")
        if item:
            resolved = self._resolve_invoice_path(item, pool="pending")
            if resolved:
                self._open_path(resolved)
                return
            delete_it = messagebox.askyesno("文件缺失", f"文件不存在：{item.stored_filename}\n是否删除这条待审批记录？")
            if delete_it:
                self._delete_pending_records([item.id], remove_files=False)
                self.refresh_all()

    def open_selected_pending_folder(self) -> None:
        ids = self._get_selected_pending_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择待审批记录")
            return
        item = self.store.get_invoice(ids[0], pool="pending")
        if item:
            resolved = self._resolve_invoice_path(item, pool="pending")
            folder = resolved.parent if resolved else self.pending_invoice_dir
            subprocess.Popen(["explorer", str(folder)])

    def clean_missing_pending_records(self) -> None:
        selected_ids = self._get_selected_pending_ids()
        target_rows = self.store.get_invoices(selected_ids, pool="pending") if selected_ids else self.store.get_pending_invoices()
        missing_ids: list[int] = []
        for row in target_rows:
            if self._resolve_invoice_path(row, pool="pending", update_store=True) is None:
                missing_ids.append(row.id)
        if not missing_ids:
            messagebox.showinfo("Info", "No missing-file records found")
            return
        if not messagebox.askyesno("Confirm Cleanup", f"Delete {len(missing_ids)} pending records whose files are missing?"):
            return
        self._delete_pending_records(missing_ids, remove_files=False)
        self.refresh_all()
        self.set_status(f"Cleaned {len(missing_ids)} invalid pending records")

    def delete_selected_pending_records(self) -> None:
        ids = self._get_selected_pending_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择待审批记录")
            return
        if not messagebox.askyesno("确认删除", f"删除选中的 {len(ids)} 条待审批记录，并尝试删除对应待审批文件？"):
            return
        deleted_count, errors = self._delete_pending_records(ids, remove_files=True)
        self.refresh_all()
        if errors:
            summary = [f"已删除 {deleted_count} 条待审批记录。", "以下文件删除失败，对应记录已保留："]
            summary.extend(errors[:12])
            messagebox.showwarning("部分删除失败", "\n".join(summary))
        else:
            messagebox.showinfo("删除完成", f"已删除 {deleted_count} 条待审批记录")
        self.set_status(f"已删除 {deleted_count} 条待审批记录")

    def batch_edit_parties(self, party_type: str) -> None:
        self.batch_edit_invoice_fields(default_field=party_type)

    def batch_edit_invoice_fields(self, default_field: str = "seller") -> None:
        ids = self._get_selected_invoice_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择发票记录")
            return
        field_options = [
            ("seller", "销售方"),
            ("buyer", "购买方"),
            ("invoice_no", "发票号码"),
            ("amount", "金额"),
            ("date", "开票日期"),
        ]
        label_by_field = dict(field_options)
        if default_field not in label_by_field:
            default_field = "seller"
        dialog = tk.Toplevel(self)
        dialog.title("批量修改字段")
        dialog.geometry("520x220")
        dialog.transient(self)
        dialog.grab_set()
        field_var = tk.StringVar(value=label_by_field[default_field])
        value_var = tk.StringVar()

        body = ttk.Frame(dialog, padding=12)
        body.pack(fill=tk.BOTH, expand=True)
        body.columnconfigure(1, weight=1)

        ttk.Label(body, text="修改字段").grid(row=0, column=0, sticky="w", pady=(0, 8))
        field_combo = ttk.Combobox(body, textvariable=field_var, values=[label for _, label in field_options], state="readonly", width=18)
        field_combo.grid(row=0, column=1, sticky="w", pady=(0, 8))

        ttk.Label(body, text="新值").grid(row=1, column=0, sticky="w")
        value_combo = ttk.Combobox(body, textvariable=value_var, width=48)
        value_combo.grid(row=1, column=1, sticky="ew")

        help_var = tk.StringVar(value="")
        ttk.Label(body, textvariable=help_var, foreground="#6E7781").grid(row=2, column=1, sticky="w", pady=(8, 0))

        def selected_field() -> str:
            current_label = field_var.get()
            for key, label in field_options:
                if label == current_label:
                    return key
            return "seller"

        def current_values() -> list[str]:
            field = selected_field()
            if field == "seller":
                return self.store.get_all_party_names("seller")
            if field == "buyer":
                return self.store.get_all_party_names("buyer")
            return []

        def refresh_value_input() -> None:
            field = selected_field()
            value_var.set("")
            if field in {"seller", "buyer"}:
                value_combo.configure(values=current_values())
                help_var.set("支持输入关键字模糊匹配白名单。")
            elif field == "date":
                value_combo.configure(values=[])
                help_var.set("日期格式：YYYY-MM-DD，例如 2026-04-27。")
            elif field == "amount":
                value_combo.configure(values=[])
                help_var.set("金额只填数字，例如 1280.50。")
            else:
                value_combo.configure(values=[])
                help_var.set("发票号码会自动去除空格并统一校验。")

        field_combo.bind("<<ComboboxSelected>>", lambda _e: refresh_value_input())
        self._bind_searchable_combobox(value_combo, current_values, value_var)
        refresh_value_input()

        def on_apply() -> None:
            field = selected_field()
            raw_value = value_var.get().strip()
            rows = self.store.get_invoices(ids, pool="approved")
            updated = 0
            skipped: list[str] = []
            for row in rows:
                seller_name = row.seller_name
                buyer_name = row.buyer_name
                invoice_number = row.invoice_number
                amount = row.amount
                invoice_date = row.invoice_date
                if field == "seller":
                    seller_name = clean_company_name(raw_value)
                    if not is_valid_company_name(seller_name):
                        messagebox.showwarning("提示", "销售方名称无效")
                        return
                elif field == "buyer":
                    buyer_name = clean_company_name(raw_value)
                    if not is_valid_company_name(buyer_name):
                        messagebox.showwarning("提示", "购买方名称无效")
                        return
                elif field == "invoice_no":
                    invoice_number = normalize_invoice_number(raw_value)
                    if not invoice_number:
                        messagebox.showwarning("提示", "发票号码无效")
                        return
                elif field == "amount":
                    try:
                        amount = float(raw_value.replace(",", "").replace("¥", "").strip())
                    except ValueError:
                        messagebox.showwarning("提示", "金额格式无效")
                        return
                elif field == "date":
                    invoice_date = raw_value
                    if not self._is_valid_date(invoice_date):
                        messagebox.showwarning("提示", "开票日期格式请使用 YYYY-MM-DD")
                        return

                if normalize_company(seller_name) == normalize_company(buyer_name):
                    skipped.append(row.original_filename or row.stored_filename)
                    continue

                if self.store.invoice_exists(seller_name, buyer_name, invoice_number, pool="approved", exclude_pool="approved", exclude_id=row.id):
                    skipped.append(f"{row.original_filename or row.stored_filename}：重复发票")
                    continue

                self.store.update_invoice_core(
                    invoice_id=row.id,
                    seller_name=seller_name,
                    buyer_name=buyer_name,
                    invoice_number=invoice_number,
                    amount=amount,
                    invoice_date=invoice_date,
                    needs_review=row.needs_review,
                    review_reason=row.review_reason,
                    approval_status=row.approval_status,
                    approval_reason=row.approval_reason,
                    pool="approved",
                )
                updated += 1
            if updated:
                for updated_item in self.store.get_invoices(ids, pool="approved"):
                    self.finalize_invoice_file(updated_item)
            dialog.destroy()
            self.refresh_all()
            message = f"已修改 {updated} 条记录"
            if skipped:
                message += f"；跳过 {len(skipped)} 条"
                messagebox.showwarning("部分跳过", "\n".join([message, *skipped[:12]]))
            self.set_status(message)

        btns = ttk.Frame(dialog)
        btns.pack(fill=tk.X, padx=12, pady=12)
        ttk.Button(btns, text="取消", command=dialog.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="应用", command=on_apply, style="Primary.TButton").pack(side=tk.RIGHT, padx=(0, 6))

    def delete_selected_invoices(self) -> None:
        ids = self._get_selected_invoice_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择发票记录")
            return
        if not messagebox.askyesno("确认删除", f"确定删除选中的 {len(ids)} 条正式发票记录吗？"):
            return
        rows = self.store.get_invoices(ids, pool="approved")
        for row in rows:
            path = Path(row.source_path)
            if path.exists():
                path.unlink()
        self.store.delete_invoices(ids, pool="approved")
        self.refresh_all()
        self.set_status(f"已删除 {len(ids)} 条正式发票记录")

    def confirm_current_pending(self) -> None:
        ids = self._get_selected_pending_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择待审批记录")
            return
        self._commit_combobox_selection(self.cbo_pending_seller, self.var_pending_seller)
        self._commit_combobox_selection(self.cbo_pending_buyer, self.var_pending_buyer)
        self._approve_records(
            ids=[ids[0]],
            seller_override=self.cbo_pending_seller.get() or self.var_pending_seller.get(),
            buyer_override=self.cbo_pending_buyer.get() or self.var_pending_buyer.get(),
            add_seller_whitelist=self.var_pending_seller_sync.get() == "whitelist",
            add_buyer_whitelist=self.var_pending_buyer_sync.get() == "whitelist",
            remember_seller_alias=self.var_pending_seller_sync.get() == "alias",
            remember_buyer_alias=self.var_pending_buyer_sync.get() == "alias",
        )

    def bulk_approve_selected(self) -> None:
        ids = self._get_selected_pending_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择待审批记录")
            return
        if not messagebox.askyesno("批量审批", f"将按现有白名单/别名核对并确认选中的 {len(ids)} 条记录。\n未命中白名单、重复或字段无效的记录会保留在审批池。是否继续？"):
            return
        self._approve_records(
            ids=ids,
            require_whitelist_match=True,
        )

    def bulk_replace_pending_party(self, party_type: str) -> None:
        ids = self._get_selected_pending_ids()
        if not ids:
            messagebox.showwarning("提示", "请先选择待审批记录")
            return
        names = self.store.get_all_party_names(party_type)
        prompt = "销售方" if party_type == "seller" else "购买方"
        dialog = tk.Toplevel(self)
        dialog.title(f"批量替换{prompt}")
        dialog.geometry("460x180")
        dialog.transient(self)
        dialog.grab_set()
        value_var = tk.StringVar()
        add_white_var = tk.BooleanVar(value=False)
        ttk.Label(dialog, text=f"替换后的{prompt}").pack(anchor="w", padx=12, pady=(12, 6))
        combo = ttk.Combobox(dialog, textvariable=value_var, values=names, width=40)
        combo.pack(fill=tk.X, padx=12)
        self._bind_searchable_combobox(combo, lambda: names, value_var)
        ttk.Checkbutton(dialog, text="同时加入白名单", variable=add_white_var).pack(anchor="w", padx=12, pady=(10, 0))

        def on_apply() -> None:
            self._commit_combobox_selection(combo, value_var)
            name = clean_company_name(value_var.get())
            if not is_valid_company_name(name):
                messagebox.showwarning("提示", f"{prompt}名称无效")
                return
            self._approve_records(
                ids=ids,
                seller_override=name if party_type == "seller" else None,
                buyer_override=name if party_type == "buyer" else None,
                add_seller_whitelist=add_white_var.get() if party_type == "seller" else False,
                add_buyer_whitelist=add_white_var.get() if party_type == "buyer" else False,
            )
            dialog.destroy()

        btns = ttk.Frame(dialog)
        btns.pack(fill=tk.X, padx=12, pady=12)
        ttk.Button(btns, text="取消", command=dialog.destroy).pack(side=tk.RIGHT)
        ttk.Button(btns, text="应用并确认", command=on_apply, style="Primary.TButton").pack(side=tk.RIGHT, padx=(0, 6))

    def _approve_records(
        self,
        ids: list[int],
        seller_override: str | None = None,
        buyer_override: str | None = None,
        add_seller_whitelist: bool = False,
        add_buyer_whitelist: bool = False,
        remember_seller_alias: bool = False,
        remember_buyer_alias: bool = False,
        require_whitelist_match: bool = False,
    ) -> None:
        rows = self.store.get_invoices(ids, pool="pending")
        if not rows:
            return
        updated = 0
        issues: list[str] = []
        duplicate_ids: list[int] = []
        for row in rows:
            target_seller = clean_company_name(seller_override or row.seller_name)
            target_buyer = clean_company_name(buyer_override or row.buyer_name)
            if not is_valid_company_name(target_seller) or not is_valid_company_name(target_buyer):
                issues.append(f"{row.original_filename} -> 主体名称无效")
                continue

            standard_seller = self.store.resolve_standard_name("seller", target_seller)
            standard_buyer = self.store.resolve_standard_name("buyer", target_buyer)
            if require_whitelist_match:
                seller_known = normalize_company(target_seller) in self.store.get_whitelist_map("seller")
                buyer_known = normalize_company(target_buyer) in self.store.get_whitelist_map("buyer")
                if not seller_known or not buyer_known:
                    missing = []
                    if not seller_known:
                        missing.append("销售方未命中白名单")
                    if not buyer_known:
                        missing.append("购买方未命中白名单")
                    issues.append(f"{row.original_filename} -> {'；'.join(missing)}")
                    continue
            if normalize_company(standard_seller) == normalize_company(standard_buyer):
                issues.append(f"{row.original_filename} -> 销售方与购买方相同")
                continue

            if add_seller_whitelist:
                self._ensure_whitelist_entry("seller", standard_seller)
            if add_buyer_whitelist:
                self._ensure_whitelist_entry("buyer", standard_buyer)

            if self.store.invoice_exists(standard_seller, standard_buyer, row.invoice_number, pool="all", exclude_pool="pending", exclude_id=row.id):
                duplicate_ids.append(row.id)
                issues.append(f"{row.original_filename} -> 重复发票：{standard_seller}_{standard_buyer}_{row.invoice_number}")
                continue

            pending_path = self._resolve_invoice_path(row, pool="pending")
            if pending_path is None or not pending_path.exists():
                issues.append(f"{row.original_filename} -> 源文件不存在")
                continue

            if remember_seller_alias:
                ok, message = self._remember_alias("seller", standard_seller, row.seller_name)
                if not ok:
                    issues.append(f"{row.original_filename} -> 销售方别名未保存：{message}")
                    continue
            if remember_buyer_alias:
                ok, message = self._remember_alias("buyer", standard_buyer, row.buyer_name)
                if not ok:
                    issues.append(f"{row.original_filename} -> 购买方别名未保存：{message}")
                    continue

            final_name = self._make_final_filename(standard_seller, standard_buyer, row.invoice_number)
            target_path = self._next_available_file(self.approved_invoice_dir / final_name)
            target_path.parent.mkdir(parents=True, exist_ok=True)
            pending_path.replace(target_path)
            self.store.approve_pending_invoice(
                pending_id=row.id,
                seller_name=standard_seller,
                buyer_name=standard_buyer,
                source_path=str(target_path),
                stored_filename=target_path.name,
            )
            updated += 1

        if duplicate_ids:
            detail = "\n".join(issues[:8])
            if messagebox.askyesno(
                "发现重复发票",
                f"发现 {len(duplicate_ids)} 条重复发票，已保留在审批池。\n\n{detail}\n\n是否直接从审批池删除这些重复发票及待审批文件？",
            ):
                deleted_count, delete_errors = self._delete_pending_records(duplicate_ids, remove_files=True)
                issues.append(f"已删除重复发票 {deleted_count} 条")
                if delete_errors:
                    issues.extend([f"删除文件失败：{item}" for item in delete_errors[:5]])

        self.refresh_all()
        summary = [f"成功确认 {updated} 条"]
        if issues:
            summary.append("未处理记录：")
            summary.extend(issues[:10])
        messagebox.showinfo("审批结果", "\n".join(summary))

    def finalize_invoice_file(self, row: InvoiceRecord) -> None:
        if not is_valid_company_name(row.seller_name) or not is_valid_company_name(row.buyer_name) or not row.invoice_number:
            return
        current_path = Path(row.source_path)
        if not current_path.exists():
            return
        final_name = self._make_final_filename(row.seller_name, row.buyer_name, row.invoice_number)
        target_path = self._next_available_file(self.approved_invoice_dir / final_name)
        if current_path.resolve() == target_path.resolve():
            self.store.update_invoice_file(row.id, str(current_path), current_path.name, 1, pool="approved")
            return
        target_path.parent.mkdir(parents=True, exist_ok=True)
        current_path.replace(target_path)
        self.store.update_invoice_file(row.id, str(target_path), target_path.name, 1, pool="approved")

    def _ensure_whitelist_entry(self, party_type: str, standard_name: str) -> WhitelistEntry | None:
        cleaned = clean_company_name(standard_name)
        if not is_valid_company_name(cleaned):
            return None
        self.store.add_whitelist_item(party_type, cleaned)
        return self.store.get_whitelist_entry_by_norm(party_type, normalize_company(cleaned))

    def _remember_alias(self, party_type: str, standard_name: str, raw_name: str) -> tuple[bool, str]:
        cleaned_raw = clean_company_name(raw_name)
        cleaned_standard = clean_company_name(standard_name)
        if not is_valid_company_name(cleaned_raw) or not is_valid_company_name(cleaned_standard):
            return False, "识别值或标准名无效"
        standard_norm = normalize_company(cleaned_standard)
        raw_norm = normalize_company(cleaned_raw)
        if raw_norm == standard_norm:
            return True, ""

        entry = self.store.get_whitelist_entry_by_norm(party_type, standard_norm)
        if not entry:
            return False, "标准名不在白名单中，不能挂别名"

        mapped_name = self.store.resolve_standard_name(party_type, cleaned_raw)
        mapped_norm = normalize_company(mapped_name)
        if mapped_norm == standard_norm:
            return True, ""
        if mapped_norm != raw_norm:
            return False, f"该别名已映射到：{mapped_name}"

        ok, message = self.store.add_whitelist_alias(entry.id, cleaned_raw)
        return ok, message

    def _get_filtered_invoice_rows(self) -> list[InvoiceRecord]:
        seller, buyer, start_month, end_month, keyword = self.get_invoice_filters()
        return self.store.query_all_filtered(
            seller,
            buyer,
            start_month,
            end_month,
            "",
            keyword,
            self.sort_column,
            self.sort_ascending,
            pool="approved",
        )

    def export_summary_excel(self) -> None:
        rows = self._get_filtered_invoice_rows()
        if not rows:
            messagebox.showinfo("提示", "当前筛选条件下没有可导出的记录")
            return
        file_path = filedialog.asksaveasfilename(
            title="导出汇总表",
            defaultextension=".xlsx",
            initialfile=f"发票汇总表_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not file_path:
            return
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "发票汇总"
        headers = ["销售方", "购买方", "发票号码", "金额", "开票日期", "文件名", "导入时间"]
        sheet.append(headers)
        for row in rows:
            sheet.append(
                [
                    row.seller_name,
                    row.buyer_name,
                    row.invoice_number,
                    row.amount,
                    row.invoice_date,
                    row.stored_filename,
                    row.imported_at,
                ]
            )
        sheet.freeze_panes = "A2"
        for col, width in {"A": 26, "B": 26, "C": 22, "D": 14, "E": 14, "F": 40, "G": 22}.items():
            sheet.column_dimensions[col].width = width
        workbook.save(file_path)
        messagebox.showinfo("导出完成", f"已导出汇总表：\n{file_path}")

    def export_filtered_pdfs(self) -> None:
        rows = self._get_filtered_invoice_rows()
        if not rows:
            messagebox.showinfo("提示", "当前筛选条件下没有可导出的记录")
            return
        target_dir = filedialog.askdirectory(title="选择PDF导出目录")
        if not target_dir:
            return
        out_dir = Path(target_dir) / f"发票PDF_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        out_dir.mkdir(parents=True, exist_ok=True)
        exported = 0
        for row in rows:
            src = self._resolve_invoice_path(row, pool="approved")
            if src and src.exists():
                shutil.copy2(src, self._next_available_file(out_dir / src.name))
                exported += 1
        messagebox.showinfo("导出完成", f"已导出 PDF {exported} 份\n目录：{out_dir}")

    def export_filtered_grouped(self) -> None:
        rows = self._get_filtered_invoice_rows()
        if not rows:
            messagebox.showinfo("提示", "当前筛选条件下没有可导出的记录")
            return
        target_dir = filedialog.askdirectory(title="选择导出根目录")
        if not target_dir:
            return
        root_dir = Path(target_dir) / f"发票PDF_分目录_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        root_dir.mkdir(parents=True, exist_ok=True)
        for row in rows:
            src = self._resolve_invoice_path(row, pool="approved")
            if not src or not src.exists():
                continue
            month_folder = self._resolve_month_folder(row.invoice_date)
            seller_folder = sanitize_filename(row.seller_name or "未知销售方")
            out_dir = root_dir / month_folder / seller_folder
            out_dir.mkdir(parents=True, exist_ok=True)
            shutil.copy2(src, self._next_available_file(out_dir / src.name))
        messagebox.showinfo("导出完成", f"已按 月份/销售方 导出 {len(rows)} 条记录\n目录：{root_dir}")

    def _resolve_month_folder(self, invoice_date: str) -> str:
        value = (invoice_date or "").strip()
        if len(value) >= 7 and value[4] == "-":
            return value[:7]
        return "未知月份"

    def refresh_whitelist_views(self) -> None:
        for party_type in ("seller", "buyer"):
            self.refresh_whitelist_view(party_type)

    def on_whitelist_select(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_entries: ttk.Treeview = widgets["entries"]  # type: ignore[assignment]
        selected = tree_entries.selection()
        whitelist_id = int(selected[0]) if selected else None
        self.refresh_whitelist_alias_view(party_type, whitelist_id)

    def refresh_whitelist_view(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_entries: ttk.Treeview = widgets["entries"]  # type: ignore[assignment]
        selected = tree_entries.selection()
        selected_id = selected[0] if selected else None
        for item_id in tree_entries.get_children():
            tree_entries.delete(item_id)

        search_text = clean_company_name(self.whitelist_search_vars[party_type].get())
        rows = self.store.get_whitelist_entries(party_type)
        if search_text:
            search_norm = normalize_company(search_text)
            filtered_rows = []
            for row in rows:
                alias_rows = self.store.get_whitelist_aliases(row.id)
                alias_text = " ".join(alias.alias_name for alias in alias_rows)
                alias_norm = normalize_company(alias_text)
                if search_norm in row.standard_norm or (alias_norm and search_norm in alias_norm):
                    filtered_rows.append(row)
            rows = filtered_rows

        sort_key = str(widgets["entry_sort_key"])
        ascending = bool(widgets["entry_sort_ascending"])
        if sort_key == "id":
            rows.sort(key=lambda row: row.id, reverse=not ascending)
        else:
            rows.sort(key=lambda row: row.standard_name, reverse=not ascending)

        for row in rows:
            tree_entries.insert("", tk.END, iid=str(row.id), values=(row.id, row.standard_name))
        if selected_id and tree_entries.exists(selected_id):
            tree_entries.selection_set(selected_id)
        self.refresh_whitelist_alias_view(party_type)

    def refresh_whitelist_alias_view(self, party_type: str, whitelist_id: int | None = None) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_entries: ttk.Treeview = widgets["entries"]  # type: ignore[assignment]
        tree_aliases: ttk.Treeview = widgets["aliases"]  # type: ignore[assignment]
        if whitelist_id is None:
            selected = tree_entries.selection()
            whitelist_id = int(selected[0]) if selected else None
        for item_id in tree_aliases.get_children():
            tree_aliases.delete(item_id)
        if not whitelist_id:
            return
        aliases = self.store.get_whitelist_aliases(whitelist_id)
        search_text = clean_company_name(self.whitelist_alias_search_vars[party_type].get())
        if search_text:
            search_norm = normalize_company(search_text)
            aliases = [alias for alias in aliases if search_norm in alias.alias_norm]
        sort_key = str(widgets["alias_sort_key"])
        ascending = bool(widgets["alias_sort_ascending"])
        if sort_key == "id":
            aliases.sort(key=lambda row: row.id, reverse=not ascending)
        else:
            aliases.sort(key=lambda row: row.alias_name, reverse=not ascending)
        for alias in aliases:
            tree_aliases.insert("", tk.END, iid=str(alias.id), values=(alias.id, alias.alias_name))

    def sort_whitelist_entries(self, party_type: str, key: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        current_key = str(widgets["entry_sort_key"])
        current_ascending = bool(widgets["entry_sort_ascending"])
        widgets["entry_sort_key"] = key
        widgets["entry_sort_ascending"] = (not current_ascending) if current_key == key else True
        self.refresh_whitelist_view(party_type)

    def sort_whitelist_aliases(self, party_type: str, key: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        current_key = str(widgets["alias_sort_key"])
        current_ascending = bool(widgets["alias_sort_ascending"])
        widgets["alias_sort_key"] = key
        widgets["alias_sort_ascending"] = (not current_ascending) if current_key == key else True
        self.refresh_whitelist_alias_view(party_type)

    def add_whitelist_entry(self, party_type: str) -> None:
        name = simpledialog.askstring("新增标准名", "请输入标准名称：", parent=self)
        if not name:
            return
        inserted, existing = self.store.add_whitelist_item(party_type, name)
        self.refresh_whitelist_views()
        self.refresh_filter_options()
        if inserted:
            self.set_status("白名单已添加")
        else:
            messagebox.showinfo("提示", f"已存在：{existing or name}")

    def edit_whitelist_entry(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_entries: ttk.Treeview = widgets["entries"]  # type: ignore[assignment]
        selected = tree_entries.selection()
        if not selected:
            messagebox.showwarning("提示", "请先选择标准名")
            return
        current = tree_entries.item(selected[0], "values")[1]
        new_name = simpledialog.askstring("编辑标准名", "请输入新的标准名：", initialvalue=current, parent=self)
        if not new_name:
            return
        ok, message = self.store.update_whitelist_item(int(selected[0]), new_name)
        if not ok:
            messagebox.showwarning("提示", message)
            return
        self.refresh_whitelist_views()
        self.refresh_filter_options()

    def delete_whitelist_entry(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_entries: ttk.Treeview = widgets["entries"]  # type: ignore[assignment]
        selected = tree_entries.selection()
        if not selected:
            messagebox.showwarning("提示", "请先选择标准名")
            return
        if not messagebox.askyesno("删除确认", "确定删除选中的标准名及其别名吗？"):
            return
        self.store.delete_whitelist_item(int(selected[0]))
        self.refresh_whitelist_views()
        self.refresh_filter_options()

    def add_whitelist_alias(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_entries: ttk.Treeview = widgets["entries"]  # type: ignore[assignment]
        selected = tree_entries.selection()
        if not selected:
            messagebox.showwarning("提示", "请先选择标准名")
            return
        alias = simpledialog.askstring("新增别名", "请输入别名：", parent=self)
        if not alias:
            return
        ok, message = self.store.add_whitelist_alias(int(selected[0]), alias)
        if not ok:
            messagebox.showwarning("提示", message)
            return
        self.on_whitelist_select(party_type)

    def edit_whitelist_alias(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_aliases: ttk.Treeview = widgets["aliases"]  # type: ignore[assignment]
        selected = tree_aliases.selection()
        if not selected:
            messagebox.showwarning("提示", "请先选择别名")
            return
        current = tree_aliases.item(selected[0], "values")[1]
        new_name = simpledialog.askstring("编辑别名", "请输入新的别名：", initialvalue=current, parent=self)
        if not new_name:
            return
        ok, message = self.store.update_whitelist_alias(int(selected[0]), new_name)
        if not ok:
            messagebox.showwarning("提示", message)
            return
        self.on_whitelist_select(party_type)

    def delete_whitelist_alias(self, party_type: str) -> None:
        widgets = self.whitelist_widgets[party_type]
        tree_aliases: ttk.Treeview = widgets["aliases"]  # type: ignore[assignment]
        selected = tree_aliases.selection()
        if not selected:
            messagebox.showwarning("提示", "请先选择别名")
            return
        self.store.delete_whitelist_alias(int(selected[0]))
        self.on_whitelist_select(party_type)

    def import_whitelist_excel(self, party_type: str) -> None:
        file_path = filedialog.askopenfilename(title="选择Excel文件", filetypes=[("Excel", "*.xlsx;*.xls")])
        if not file_path:
            return
        try:
            workbook = load_workbook(file_path, data_only=True)
            sheet = workbook.active
            names: list[str] = []
            for row in sheet.iter_rows(min_row=1, max_col=1, values_only=True):
                value = row[0]
                if value is not None:
                    text = str(value).strip()
                    if text:
                        names.append(text)
            inserted, duplicates = self.store.add_whitelist_items(party_type, names)
            self.refresh_whitelist_views()
            self.refresh_filter_options()
            message_lines = [f"成功导入 {inserted} 条"]
            if duplicates:
                message_lines.append(f"重复 {len(duplicates)} 条：")
                message_lines.extend(duplicates[:10])
            messagebox.showinfo("导入完成", "\n".join(message_lines))
        except Exception as exc:
            messagebox.showerror("导入失败", str(exc))

    def download_reconcile_template(self) -> None:
        file_path = filedialog.asksaveasfilename(
            title="保存对账模板",
            defaultextension=".xlsx",
            initialfile=f"批量对账模板_{datetime.now().strftime('%Y%m%d')}.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not file_path:
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "对账清单"
        headers = ["销售方", "购买方", "发票号", "金额"]
        sheet.append(headers)
        sheet.append(["", "", "", ""])
        sheet.freeze_panes = "A2"
        widths = {
            "A": 28,
            "B": 28,
            "C": 22,
            "D": 14,
        }
        for col, width in widths.items():
            sheet.column_dimensions[col].width = width
        workbook.save(file_path)
        self.set_status("已生成批量对账模板")
        messagebox.showinfo("模板已生成", "批量对账模板已保存。")

    def import_reconcile_sheet(self) -> None:
        file_path = filedialog.askopenfilename(
            title="选择对账清单",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("CSV 文件", "*.csv")],
        )
        if not file_path:
            return
        try:
            input_rows = self._read_reconcile_input_rows(Path(file_path))
            if not input_rows:
                messagebox.showinfo("导入结果", "对账清单中没有可用数据。")
                return
            self.batch_reconcile_results = self._run_batch_reconcile(input_rows)
            self._populate_batch_reconcile_tree()
            summary = self._summarize_batch_reconcile(self.batch_reconcile_results)
            self.set_status(summary)
            messagebox.showinfo("导入完成", summary)
        except Exception as exc:
            messagebox.showerror("导入失败", str(exc))

    def export_reconcile_results(self) -> None:
        if not self.batch_reconcile_results:
            messagebox.showwarning("提示", "请先导入并执行批量对账。")
            return
        file_path = filedialog.asksaveasfilename(
            title="导出对账结果",
            defaultextension=".xlsx",
            initialfile=f"批量对账结果_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx"), ("CSV 文件", "*.csv")],
        )
        if not file_path:
            return
        headers = [
            ("input_seller", "输入销售方"),
            ("input_buyer", "输入购买方"),
            ("input_invoice_no", "输入发票号"),
            ("input_amount", "输入金额"),
            ("result", "核对结果"),
            ("matched_seller", "系统销售方"),
            ("matched_buyer", "系统购买方"),
            ("matched_amount", "系统金额"),
            ("matched_date", "开票日期"),
            ("file", "文件名"),
        ]
        path = Path(file_path)
        if path.suffix.lower() == ".csv":
            with path.open("w", newline="", encoding="utf-8-sig") as handle:
                writer = csv.writer(handle)
                writer.writerow([label for _, label in headers])
                for row in self.batch_reconcile_results:
                    writer.writerow([row.get(key, "") for key, _ in headers])
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "对账结果"
            sheet.append([label for _, label in headers])
            for row in self.batch_reconcile_results:
                sheet.append([row.get(key, "") for key, _ in headers])
            sheet.freeze_panes = "A2"
            for col, width in {
                "A": 24,
                "B": 24,
                "C": 20,
                "D": 14,
                "E": 18,
                "F": 24,
                "G": 24,
                "H": 14,
                "I": 14,
                "J": 36,
            }.items():
                sheet.column_dimensions[col].width = width
            workbook.save(path)
        self.set_status("已导出批量对账结果")
        messagebox.showinfo("导出完成", "批量对账结果已导出。")

    def _read_reconcile_input_rows(self, file_path: Path) -> list[dict[str, str]]:
        rows: list[dict[str, str]] = []
        header_aliases = {
            "销售方": "input_seller",
            "卖方": "input_seller",
            "seller": "input_seller",
            "购买方": "input_buyer",
            "买方": "input_buyer",
            "buyer": "input_buyer",
            "发票号": "input_invoice_no",
            "发票号码": "input_invoice_no",
            "invoice_no": "input_invoice_no",
            "invoice": "input_invoice_no",
            "金额": "input_amount",
            "amount": "input_amount",
        }

        def map_headers(raw_headers: list[str]) -> list[str]:
            mapped: list[str] = []
            for header in raw_headers:
                normalized = header.strip().lower().replace(" ", "").replace("_", "")
                matched = header_aliases.get(header.strip())
                if not matched:
                    matched = header_aliases.get(normalized)
                mapped.append(matched or "")
            return mapped

        if file_path.suffix.lower() == ".csv":
            with file_path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.reader(handle)
                raw_headers = next(reader, [])
                mapped_headers = map_headers([str(item or "") for item in raw_headers])
                for index, raw_row in enumerate(reader, start=2):
                    row_data = {"input_seller": "", "input_buyer": "", "input_invoice_no": "", "input_amount": "", "source_row": str(index)}
                    for column_index, value in enumerate(raw_row):
                        if column_index >= len(mapped_headers):
                            continue
                        key = mapped_headers[column_index]
                        if key:
                            row_data[key] = str(value or "").strip()
                    if any(row_data[key] for key in ("input_seller", "input_buyer", "input_invoice_no", "input_amount")):
                        rows.append(row_data)
            return rows

        workbook = load_workbook(file_path, data_only=True)
        sheet = workbook.active
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
        mapped_headers = map_headers([str(item or "") for item in header_cells])
        for row_index, raw_row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            row_data = {"input_seller": "", "input_buyer": "", "input_invoice_no": "", "input_amount": "", "source_row": str(row_index)}
            for column_index, value in enumerate(raw_row):
                if column_index >= len(mapped_headers):
                    continue
                key = mapped_headers[column_index]
                if key:
                    row_data[key] = str(value or "").strip()
            if any(row_data[key] for key in ("input_seller", "input_buyer", "input_invoice_no", "input_amount")):
                rows.append(row_data)
        return rows

    def _run_batch_reconcile(self, input_rows: list[dict[str, str]]) -> list[dict[str, str]]:
        approved_rows = self.store.query_all_filtered(pool="approved")
        rows_by_invoice: dict[str, list[InvoiceRecord]] = {}
        for row in approved_rows:
            rows_by_invoice.setdefault(row.invoice_number_norm, []).append(row)

        results: list[dict[str, str]] = []
        for input_row in input_rows:
            seller_text = clean_company_name(input_row.get("input_seller", ""))
            buyer_text = clean_company_name(input_row.get("input_buyer", ""))
            invoice_text = (input_row.get("input_invoice_no", "") or "").strip()
            amount_text = (input_row.get("input_amount", "") or "").strip()
            invoice_norm = normalize_invoice_number(invoice_text)

            matched_row: InvoiceRecord | None = None
            result_messages: list[str] = []
            if not invoice_norm:
                result_messages.append("缺少发票号")
            else:
                candidates = rows_by_invoice.get(invoice_norm, [])
                if len(candidates) > 1:
                    narrowed = candidates
                    if seller_text:
                        seller_norm = normalize_company(seller_text)
                        narrowed = [row for row in narrowed if row.seller_name_norm == seller_norm]
                    if buyer_text:
                        buyer_norm = normalize_company(buyer_text)
                        narrowed = [row for row in narrowed if row.buyer_name_norm == buyer_norm]
                    if len(narrowed) == 1:
                        matched_row = narrowed[0]
                    else:
                        result_messages.append("多条匹配")
                elif not candidates:
                    result_messages.append("未找到")
                else:
                    matched_row = candidates[0]

            if matched_row is not None:
                if seller_text and matched_row.seller_name_norm != normalize_company(seller_text):
                    result_messages.append("销售方不一致")
                if buyer_text and matched_row.buyer_name_norm != normalize_company(buyer_text):
                    result_messages.append("购买方不一致")
                if amount_text:
                    try:
                        target_amount = float(amount_text.replace(",", ""))
                    except ValueError:
                        result_messages.append("输入金额格式错误")
                    else:
                        if abs(matched_row.amount - target_amount) > 0.009:
                            result_messages.append("金额不一致")

            result_text = "匹配" if not result_messages else "；".join(dict.fromkeys(result_messages))
            results.append(
                {
                    "source_row": input_row.get("source_row", ""),
                    "input_seller": seller_text,
                    "input_buyer": buyer_text,
                    "input_invoice_no": invoice_text,
                    "input_amount": amount_text,
                    "result": result_text,
                    "matched_seller": matched_row.seller_name if matched_row else "",
                    "matched_buyer": matched_row.buyer_name if matched_row else "",
                    "matched_amount": f"{matched_row.amount:.2f}" if matched_row else "",
                    "matched_date": matched_row.invoice_date if matched_row else "",
                    "file": matched_row.stored_filename if matched_row else "",
                }
            )
        return results

    def _populate_batch_reconcile_tree(self) -> None:
        for item_id in self.tree_batch_reconcile.get_children():
            self.tree_batch_reconcile.delete(item_id)
        only_issues = self.var_batch_reconcile_only_issues.get()
        visible_rows = self.batch_reconcile_results
        if only_issues:
            visible_rows = [row for row in self.batch_reconcile_results if row.get("result", "") != "匹配"]
        key = self.batch_reconcile_sort_column
        ascending = self.batch_reconcile_sort_ascending

        def sort_value(row: dict[str, str]) -> object:
            value = row.get(key, "")
            if key in {"input_amount", "matched_amount"}:
                try:
                    return float(value or 0)
                except ValueError:
                    return -1.0
            return value or ""

        visible_rows = sorted(visible_rows, key=sort_value, reverse=not ascending)
        for row in visible_rows:
            self.tree_batch_reconcile.insert(
                "",
                tk.END,
                iid=row.get("source_row", ""),
                values=(
                    row.get("input_seller", ""),
                    row.get("input_buyer", ""),
                    row.get("input_invoice_no", ""),
                    row.get("input_amount", ""),
                    row.get("result", ""),
                    row.get("matched_seller", ""),
                    row.get("matched_buyer", ""),
                    row.get("matched_amount", ""),
                    row.get("matched_date", ""),
                    row.get("file", ""),
                ),
            )

    def open_selected_batch_reconcile_file(self) -> None:
        selected = self.tree_batch_reconcile.selection()
        if not selected:
            return
        row_id = selected[0]
        target = next((item for item in self.batch_reconcile_results if item.get("source_row", "") == row_id), None)
        if not target:
            return
        file_name = target.get("file", "")
        if not file_name:
            messagebox.showinfo("提示", "该行没有匹配到本地发票文件。")
            return
        file_path = self.approved_invoice_dir / file_name
        if not file_path.exists():
            messagebox.showwarning("提示", f"文件不存在：{file_path}")
            return
        self._open_path(file_path)

    @staticmethod
    def _summarize_batch_reconcile(results: list[dict[str, str]]) -> str:
        counters: dict[str, int] = {}
        for row in results:
            key = row.get("result", "未分类")
            counters[key] = counters.get(key, 0) + 1
        ordered = sorted(counters.items(), key=lambda item: item[0])
        summary_parts = [f"{label} {count} 条" for label, count in ordered]
        return f"批量对账完成，共 {len(results)} 条；" + "，".join(summary_parts)

    def run_reconcile(self) -> None:
        seller_keyword = self.var_rec_seller.get().strip()
        buyer_keyword = self.var_rec_buyer.get().strip()
        invoice_no = self.var_rec_no.get().strip()
        amount_text = self.var_rec_amount.get().strip()
        try:
            target_amount = float(amount_text) if amount_text else None
        except ValueError:
            messagebox.showwarning("格式提示", "金额请输入数字")
            return
        rows = self.store.query_all_filtered(seller_keyword, buyer_keyword, "", "", "", "", "date", False, pool="approved")
        if invoice_no:
            invoice_no_norm = normalize_invoice_number(invoice_no)
            rows = [row for row in rows if row.invoice_number_norm == invoice_no_norm]
        if not rows:
            self.reconcile_rows = []
            self._populate_reconcile_tree()
            messagebox.showinfo("手动核对", "当前条件下没有匹配记录")
            self.set_status("手动核对完成，共 0 条")
            return
        self.reconcile_rows = []
        for row in rows:
            status = "匹配"
            if target_amount is not None and abs(row.amount - target_amount) > 0.009:
                status = "金额不一致"
            self.reconcile_rows.append(
                {
                    "id": row.id,
                    "seller": row.seller_name,
                    "buyer": row.buyer_name,
                    "invoice_no": row.invoice_number,
                    "amount": row.amount,
                    "date": row.invoice_date,
                    "status": status,
                    "file": row.stored_filename,
                }
            )
        self._populate_reconcile_tree()
        self.set_status(f"手动核对完成，共 {len(rows)} 条")

    def _populate_reconcile_tree(self) -> None:
        for item_id in self.tree_reconcile.get_children():
            self.tree_reconcile.delete(item_id)

        key = self.reconcile_sort_column
        ascending = self.reconcile_sort_ascending

        def sort_value(row: dict[str, object]) -> object:
            value = row.get(key, "")
            if key == "amount":
                return float(value or 0)
            return value or ""

        for row in sorted(self.reconcile_rows, key=sort_value, reverse=not ascending):
            self.tree_reconcile.insert(
                "",
                tk.END,
                iid=str(row["id"]),
                values=(
                    row["seller"],
                    row["buyer"],
                    row["invoice_no"],
                    f"{float(row['amount']):.2f}",
                    row["date"],
                    row["status"],
                    row["file"],
                ),
            )

    def open_selected_reconcile_file(self) -> None:
        selected = self.tree_reconcile.selection()
        if not selected:
            return
        item = self.store.get_invoice(int(selected[0]), pool="approved")
        if not item:
            return
        resolved = self._resolve_invoice_path(item, pool="approved")
        if resolved:
            self._open_path(resolved)
        else:
            messagebox.showwarning("提示", f"文件不存在：{item.stored_filename}")

    def show_month_picker(self, target: str) -> None:
        self.month_picker_target = target
        if self.month_picker_popup is not None and self.month_picker_popup.winfo_exists():
            self._render_month_picker()
            self.month_picker_popup.lift()
            return
        current_var = self.var_start_month if target == "start" else self.var_end_month
        current = self.normalize_month_text(current_var.get().strip())
        if current and self._is_valid_month(current):
            self.month_picker_year = int(current.split("-")[0])
        elif self.month_options:
            self.month_picker_year = int(self.month_options[0].split("-")[0])
        else:
            self.month_picker_year = datetime.now().year
        popup = tk.Toplevel(self)
        popup.title("选择月份")
        popup.transient(self)
        popup.resizable(False, False)
        popup.configure(bg="white")
        popup.protocol("WM_DELETE_WINDOW", self.close_month_picker)
        popup.bind("<Escape>", lambda _e: self.close_month_picker())
        self.month_picker_popup = popup
        self._render_month_picker()
        self.update_idletasks()
        target_widget = self.cbo_start_month if target == "start" else self.cbo_end_month
        x = target_widget.winfo_rootx()
        y = target_widget.winfo_rooty() + target_widget.winfo_height() + 2
        popup.geometry(f"+{x}+{y}")

    def _render_month_picker(self) -> None:
        if self.month_picker_popup is None or not self.month_picker_popup.winfo_exists():
            return
        popup = self.month_picker_popup
        for child in popup.winfo_children():
            child.destroy()
        wrapper = ttk.Frame(popup, padding=12, style="White.TFrame")
        wrapper.pack(fill=tk.BOTH, expand=True)
        header = ttk.Frame(wrapper, style="White.TFrame")
        header.pack(fill=tk.X, pady=(0, 10))
        ttk.Button(header, text="<", width=3, command=lambda: self.change_month_picker_year(-1)).pack(side=tk.LEFT)
        picker_title = "开始月份" if self.month_picker_target == "start" else "结束月份"
        ttk.Label(header, text=f"{picker_title}  {self.month_picker_year}", style="SummaryLabel.TLabel").pack(side=tk.LEFT, padx=10)
        ttk.Button(header, text=">", width=3, command=lambda: self.change_month_picker_year(1)).pack(side=tk.LEFT)
        selected_var = self.var_start_month if self.month_picker_target == "start" else self.var_end_month
        selected = self.normalize_month_text(selected_var.get().strip())
        selected_year = selected_month = None
        if selected and self._is_valid_month(selected):
            selected_year, selected_month = [int(part) for part in selected.split("-")]
        available_months = set(self.month_options)
        month_grid = ttk.Frame(wrapper, style="White.TFrame")
        month_grid.pack(fill=tk.BOTH, expand=True)
        for month in range(1, 13):
            row = (month - 1) // 4
            col = (month - 1) % 4
            month_value = f"{self.month_picker_year:04d}-{month:02d}"
            available = month_value in available_months
            style_name = "Primary.TButton" if selected_year == self.month_picker_year and selected_month == month else "TButton"
            button = ttk.Button(
                month_grid,
                text=f"{month}月",
                width=8,
                style=style_name,
                command=lambda m=month: self.select_month_value(self.month_picker_year, m),
            )
            if not available:
                button.state(["disabled"])
            button.grid(row=row, column=col, padx=4, pady=4, sticky="nsew")
        footer = ttk.Frame(wrapper, style="White.TFrame")
        footer.pack(fill=tk.X, pady=(10, 0))
        ttk.Button(footer, text="清除", command=self.clear_month_value).pack(side=tk.LEFT)
        ttk.Button(footer, text="本月", command=self.use_current_month).pack(side=tk.RIGHT)

    def close_month_picker(self) -> None:
        if self.month_picker_popup is not None and self.month_picker_popup.winfo_exists():
            self.month_picker_popup.destroy()
        self.month_picker_popup = None

    def change_month_picker_year(self, delta: int) -> None:
        self.month_picker_year += delta
        self._render_month_picker()

    def select_month_value(self, year: int, month: int) -> None:
        target_var = self.var_start_month if self.month_picker_target == "start" else self.var_end_month
        target_var.set(f"{year:04d}-{month:02d}")
        self.close_month_picker()
        self.apply_filter(show_no_result_feedback=True)

    def clear_month_value(self) -> None:
        target_var = self.var_start_month if self.month_picker_target == "start" else self.var_end_month
        target_var.set("")
        self.close_month_picker()
        self.apply_filter(show_no_result_feedback=True)

    def use_current_month(self) -> None:
        today = datetime.now()
        self.select_month_value(today.year, today.month)

    def _sync_invoice_summary_layout(self) -> None:
        if not hasattr(self, "summary_cells"):
            return
        columns = list(self.tree_invoice["columns"])
        if "amount" not in columns:
            return
        summary_frame = self.summary_cells["frame"]
        amount_label = self.summary_cells["amount"]
        if not isinstance(summary_frame, ttk.Frame) or not isinstance(amount_label, ttk.Label):
            return

        summary_width = max(self.tree_invoice.winfo_width(), 200)
        summary_frame.configure(width=summary_width)

        widths = [int(self.tree_invoice.column(key, "width")) for key in columns]
        amount_index = columns.index("amount")
        amount_left = sum(widths[:amount_index])
        amount_width = widths[amount_index]
        total_width = sum(widths)
        visible_width = max(self.tree_invoice.winfo_width(), 1)
        xview = self.tree_invoice.xview()
        x_fraction = float(xview[0]) if xview else 0.0
        scrollable_width = max(total_width - visible_width, 0)
        offset = int(round(scrollable_width * x_fraction))
        display_left = amount_left - offset
        amount_label.place_configure(
            x=max(display_left + 6, 6),
            y=8,
            width=max(amount_width - 12, 60),
            height=24,
        )

    @staticmethod
    def normalize_month_text(value: str) -> str:
        raw = (value or "").strip()
        if not raw:
            return ""
        compact = (
            raw.replace("年", "-")
            .replace("月", "")
            .replace("/", "-")
            .replace(".", "-")
            .replace("_", "-")
            .replace(" ", "")
        )
        if compact.isdigit() and len(compact) == 6:
            compact = f"{compact[:4]}-{compact[4:]}"
        parts = [part for part in compact.split("-") if part]
        if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
            return f"{parts[0].zfill(4)}-{parts[1].zfill(2)}"
        return raw

    @staticmethod
    def _is_valid_month(value: str) -> bool:
        try:
            datetime.strptime(value, "%Y-%m")
            return True
        except ValueError:
            return False

    @staticmethod
    def _is_valid_date(value: str) -> bool:
        try:
            datetime.strptime(value, "%Y-%m-%d")
            return True
        except ValueError:
            return False


def main() -> None:
    configure_tk_runtime()
    app = InvoiceAssistantApp()
    app.mainloop()


if __name__ == "__main__":
    main()

